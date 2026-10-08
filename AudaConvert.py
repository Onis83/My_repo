#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AUDATEX Converter v9.5.3 — BRE Client Fix Edition
- Исправлен парсинг BRE Client4.html (сшивание разорванных заголовков)
- Уточнены границы блоков работ/окраски/материалов/запчастей
- Ленивая инициализация словаря (без модальных окон при старте)
- Тёмная/светлая тема
"""
__version__ = "9.5.3"
__author__ = "Полуницкий Е.В."

import sys
import os
import re
import json
import time
from typing import List, Dict, Tuple, Optional
from contextlib import contextmanager

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QTabWidget, QPlainTextEdit, QTableWidget,
    QTableWidgetItem, QHeaderView, QFileDialog, QMessageBox,
    QMenu, QInputDialog, QCheckBox, QComboBox, QTreeWidget,
    QTreeWidgetItem, QGroupBox, QAbstractItemView, QStyledItemDelegate,
    QStyleOptionViewItem, QStyle, QDialog, QLineEdit, QTextEdit,
    QSplashScreen, QProgressBar, QGridLayout, QScrollArea, QFrame
)
from PyQt6.QtCore import Qt, QPoint, QModelIndex, QTimer, QSize, QRect
from PyQt6.QtGui import (
    QFont, QColor, QAction, QPainter, QTextCursor, QTextDocument,
    QPalette, QPixmap, QIcon, QFontDatabase
)

try:
    from bs4 import BeautifulSoup
    HAS_BS4 = True
except ImportError:
    HAS_BS4 = False

try:
    from odf.opendocument import OpenDocumentSpreadsheet, OpenDocumentText
    from odf.text import P
    from odf.table import Table, TableRow, TableCell, TableColumn
    from odf.style import Style, TextProperties, ParagraphProperties
    HAS_ODF = True
except ImportError:
    HAS_ODF = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font as XlFont, Alignment, PatternFill, Border, Side
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ============================================================
# УНИВЕРСАЛЬНЫЙ ПАРСЕР AUDATEX
# ============================================================

class AudatexParser:

    @staticmethod
    def preprocess_html(raw_html: str) -> str:
        """
        Предобработка HTML BRE Client:
        1. Сшивает разорванные буквы заголовков [С][т][о][и][м][о][с][т][ь] → Стоимость
        2. Удаляет мусорные ссылки
        3. Возвращает чистый текст для парсинга
        """
        if not HAS_BS4:
            return raw_html

        soup = BeautifulSoup(raw_html, "html.parser")

        # Удаляем скрипты, стили, навигацию
        for tag in soup(["script", "style", "nav", "header", "footer", "a"]):
            tag.decompose()

        text = soup.get_text(separator="\n")

        # Сшиваем разорванные буквы: "[С][т][о][и][м][о][с][т][ь]" → "Стоимость"
        # Паттерн находит последовательности вида [X][y][z]...
        def stitch_letters(match):
            letters = re.findall(r'\[([^\]]+)\]', match.group(0))
            return ''.join(letters)

        text = re.sub(r'(?:\[[^\]]+\]\s*){2,}', stitch_letters, text)

        # Нормализуем пробелы и переносы
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]{2,}", " ", text)

        return text.strip()

    @staticmethod
    def parse(text: str) -> Dict[str, List[Dict[str, str]]]:
        result = {'works': [], 'paint': [], 'materials': [], 'parts': []}
        lines = text.split('\n')

        result['parts'] = AudatexParser._parse_parts(lines)
        result['works'] = AudatexParser._parse_works(lines)
        result['paint'] = AudatexParser._parse_paint(lines)
        result['materials'] = AudatexParser._parse_materials(lines)

        return result

    @staticmethod
    def find_extra_paint_lines(text: str) -> Dict[str, List[Dict[str, str]]]:
        lines = text.split('\n')
        result = {'works': [], 'materials': []}

        skip_keywords = [
            'СТОИМОСТЬ РАБОТ', 'ОБЩЕЕ ВРЕМЯ', 'ИТОГО', 'ВСЕГО',
            'RUR/ЧАС', 'EUR/ЧАС', 'USD/ЧАС', 'НОРМА ВРЕМЕНИ',
            'ВАЛЮТНЫЙ КУРС', 'РЕМОНТ-КАЛЬКУЛЯЦИЯ', 'ОКОНЧАТЕЛЬНАЯ',
            'СИСТЕМА AUDATEX', 'ЛИСТ',
        ]

        def should_skip(stripped_line: str) -> bool:
            return any(kw in stripped_line for kw in skip_keywords)

        def extract_number_at_end(line: str) -> Optional[str]:
            stripped = line.rstrip()
            match = re.search(r'([\d][\d\s]*[\d]|\d)\s*$', stripped)
            if match:
                return match.group(1).replace(' ', '')
            return None

        # Поиск дополнительных работ по окраске
        in_works_block = False
        for line in lines:
            stripped = line.strip()
            if 'ЗАТРАТЫ ВРЕМЕНИ НА ОКРАСКУ' in stripped:
                in_works_block = True
                continue

            if in_works_block and any(m in stripped for m in [
                'ОКРАСКА-ЗАТРАТЫ', 'ОКОНЧАТЕЛЬНАЯ', 'РЕМОНТ-КАЛЬКУЛЯЦИЯ', 'СИСТЕМА AUDATEX'
            ]):
                in_works_block = False
                continue

            if in_works_block and stripped:
                if should_skip(stripped):
                    continue

                number = extract_number_at_end(stripped)
                if number:
                    description = re.sub(
                        r'\s*[\d][\d\s]*[\d]?\s*$', '', stripped).strip()
                    description = re.sub(r'\s{2,}', ' ', description).strip()
                    if description:
                        result['works'].append(
                            {'description': description, 'rp': number})

        # Поиск дополнительных материалов по окраске
        in_materials_block = False
        for line in lines:
            stripped = line.strip()
            if 'ОКРАСКА-ЗАТРАТЫ НА МАТЕРИАЛ' in stripped:
                in_materials_block = True
                continue

            if in_materials_block and any(m in stripped for m in [
                'ОКОНЧАТЕЛЬНАЯ', 'РЕМОНТ-КАЛЬКУЛЯЦИЯ', 'СИСТЕМА AUDATEX', 'ЗАТРАТЫ ВРЕМЕНИ'
            ]):
                in_materials_block = False
                continue

            if in_materials_block and stripped:
                if should_skip(stripped):
                    continue

                number = extract_number_at_end(stripped)
                if number:
                    description = re.sub(
                        r'\s*[\d][\d\s]*[\d]?\s*$', '', stripped).strip()
                    description = re.sub(r'\s{2,}', ' ', description).strip()
                    if '%' in description:
                        continue
                    if description:
                        result['materials'].append(
                            {'description': description, 'cost': number})

        return result

    @staticmethod
    def _normalize_spaces(text: str) -> str:
        return re.sub(r'\s{2,}', ' ', text).strip()

    @staticmethod
    def _split_code_and_description(text: str) -> Tuple[str, str]:
        text = text.strip()
        if not text:
            return "", ""

        code_patterns = [
            r'^Б/Н',
            r'^\d{2}-\d{4}\s+[A-Z0-9]+',
            r'^\d{2}\s+\d{3}\s+\d{1,2}\)?',
            r'^[A-Z0-9]{10,}\)?',
            r'^\d{4,9}\)?',
        ]

        for pattern in code_patterns:
            match = re.match(pattern, text)
            if match:
                code = match.group(0).strip()
                description = text[match.end():].strip()
                description = AudatexParser._normalize_spaces(description)
                return code, description

        match = re.search(r'\s{2,}', text)
        if match:
            code = text[:match.start()].strip()
            description = text[match.end():].strip()
            description = AudatexParser._normalize_spaces(description)
            return code, description

        parts = text.split(None, 1)
        code = parts[0] if parts else ""
        description = parts[1] if len(parts) > 1 else ""
        description = AudatexParser._normalize_spaces(description)
        return code, description

    @staticmethod
    def _is_continuation_line(line: str) -> bool:
        stripped = line.strip()
        if not stripped:
            return True

        has_numbers_at_end = bool(re.search(r'\d+\*?\s+\d+\s*$', stripped))
        has_three_numbers_at_end = bool(
            re.search(r'\d+\s+\d+\*?\s+\d+\s*$', stripped))

        if has_numbers_at_end or has_three_numbers_at_end:
            return False

        if stripped.startswith(('ВКЛ:', 'НЕ ВКЛ:', 'ВКЛ :', 'НЕ ВКЛ:')):
            return True

        leading_spaces = len(line) - len(line.lstrip())
        if leading_spaces > 5:
            return True

        return False

    @staticmethod
    def _find_marker(line: str, markers: List[str]) -> bool:
        """Ищет маркер как с пробелами, так и без."""
        stripped = line.strip()
        for m in markers:
            if m in stripped:
                return True
            spaced_m = " ".join(m)
            if spaced_m in stripped:
                return True
        return False

    @staticmethod
    def _extract_section(lines: List[str], start_markers: List[str], end_markers: List[str]) -> List[str]:
        result = []
        in_section = False

        for line in lines:
            if not in_section and AudatexParser._find_marker(line, start_markers):
                in_section = True
                continue

            if in_section:
                if AudatexParser._find_marker(line, end_markers):
                    break
                result.append(line)

        return result

    @staticmethod
    def _parse_parts(lines: List[str]) -> List[Dict[str, str]]:
        result = []
        section = AudatexParser._extract_section(
            lines,
            ['З А П Ч А С Т И', 'ЗАПЧАСТИ'],
            ['СТОИМОСТЬ РАБОТ', 'О К Р А С К А', 'ОКРАСКА', 'ОКОНЧАТЕЛЬНАЯ']
        )

        for line in section:
            stripped = line.strip()
            if not stripped or stripped.startswith('-'):
                continue

            if any(h in line for h in ['УПР №', 'НАЗВАНИЕ', 'КОЛ-ВО', '№ ДЕТАЛИ', 'СТОИМ',
                                       'ИТОГО', 'ВСЕГО', 'МЕЛКИЕ ЗАПЧАСТИ', 'УРОВЕНЬ ЦЕН']):
                continue

            if AudatexParser._is_continuation_line(line):
                continue

            parts = line.split()
            if len(parts) < 3:
                continue

            try:
                cost_str = parts[-1].replace(' ', '').replace(
                    '*', '').replace('+', '').replace('>', '')
                if not cost_str.replace('.', '').isdigit():
                    continue

                remaining_parts = parts[:-1]
                remaining_str = ' '.join(remaining_parts)

                detail_match = re.search(
                    r'[+>]?([\d\s]{5,}?)\s*$', remaining_str)
                detail_number = ""
                name_with_code = remaining_str

                if detail_match:
                    potential_detail = detail_match.group(1).strip()
                    digits_count = sum(
                        1 for c in potential_detail if c.isdigit())
                    if digits_count >= 4:
                        detail_number = AudatexParser._normalize_spaces(
                            potential_detail)
                        full_match = detail_match.group(0).strip()
                        if full_match.startswith(('+', '>')):
                            detail_number = full_match[0] + detail_number
                        name_with_code = remaining_str[:detail_match.start()].strip(
                        )

                code, name = AudatexParser._split_code_and_description(
                    name_with_code)

                qty = "1"
                qty_match = re.match(
                    r'^(\d+)\s*(?:ШТ\.?\s+)?(.+)', name, re.IGNORECASE)
                if qty_match:
                    qty_val = qty_match.group(1)
                    if int(qty_val) <= 99:
                        qty = qty_val
                        name = qty_match.group(2).strip()

                if name and cost_str:
                    result.append({
                        'code': code, 'qty': qty,
                        'name': AudatexParser._normalize_spaces(name),
                        'article': detail_number, 'cost': cost_str
                    })
            except (ValueError, IndexError):
                continue

        return result

    @staticmethod
    def _parse_works(lines: List[str]) -> List[Dict[str, str]]:
        result = []
        in_section = False
        section_lines = []
        has_kl_column = False

        work_start = ['СТОИМОСТЬ РАБОТ', 'СТОИМОСТИ РАБОТ']
        work_end = ['О К Р А С К А', 'ОКРАСКА',
                    'ОКОНЧАТЕЛЬНАЯ', 'З А П Ч А С Т И', 'ЗАПЧАСТИ']

        for line in lines:
            if not in_section and any(m in line for m in work_start) and ('НОРМА ВРЕМЕНИ' in line or 'ВРЕМЕНИ' in line):
                if section_lines:
                    result.extend(AudatexParser._parse_works_lines(
                        section_lines, has_kl_column))
                in_section = True
                section_lines = []
                has_kl_column = False
                continue

            if in_section:
                if ('КЛ' in line and 'РП' in line and 'СТОИМ' in line
                        and 'РАБ.' not in line.upper()):
                    has_kl_column = True

                if AudatexParser._find_marker(line, work_end):
                    if section_lines:
                        result.extend(AudatexParser._parse_works_lines(
                            section_lines, has_kl_column))
                    in_section = False
                    section_lines = []
                    if 'ОКОНЧАТЕЛЬНАЯ' in line:
                        break
                    continue

                section_lines.append(line)

        if section_lines:
            result.extend(AudatexParser._parse_works_lines(
                section_lines, has_kl_column))

        return result

    @staticmethod
    def _parse_works_lines(lines: List[str], has_kl_column: bool = False) -> List[Dict[str, str]]:
        result = []
        for line in lines:
            stripped = line.strip()
            if not stripped or stripped.startswith('-'):
                continue

            if any(h in stripped for h in [
                '№ РАБ', 'КОД ОПЕР', 'РАБОТЫ ПО РЕМ', 'ПОЗ./', 'СИСТЕМА AUDATEX', 'ЛИСТ',
                'ИТОГО', 'ВСЕГО', 'ЗАМЕР', 'ИТОГО СТОИМОСТЬ', 'ВАЛЮТНЫЙ КУРС', 'НОРМА ВРЕМЕНИ',
                'СТОИМ/КЛ', 'RUR/ЧАС', 'СТОИМ  =', 'СТОИМ =',
                'РЕМОНТ-КАЛЬКУЛЯЦИЯ', 'РЕМОНТ - КАЛЬКУЛЯЦИЯ', 'ПОЯСНЕНИЯ', 'ПЕРЕСЧ. ЦЕНЫ',
                '* = ДАННЫЕ', 'ZAX = ЗАТРАТЫ', '(C) ВСЕ ПРАВА', 'AUTOMOTIVE GMBH',
                '№ ДЕЛА', 'ПРОИЗВОД', 'КУЗОВ №', 'ПРОБЕГ', 'ВАРИАНТЫ', 'ГОС.№',
                'КОД ТИПА', 'КОНСТРУКЦИИ', 'КОНСТРУКЦИ',
            ]):
                continue

            if re.search(r'\b\d+\s+(RUR|EUR|USD)\b', stripped):
                continue

            if AudatexParser._is_continuation_line(line):
                result.append({
                    'code': '', 'description': AudatexParser._normalize_spaces(stripped),
                    'rp': '', 'cost': ''
                })
                continue

            parts = line.split()
            if len(parts) < 3:
                continue

            try:
                cost_str = parts[-1].replace(' ', '').replace('*', '')
                rp_str = parts[-2].replace(' ',
                                           '').replace('*', '').replace(')', '')

                if not cost_str.replace('.', '').isdigit():
                    continue
                if not rp_str.replace('.', '').isdigit():
                    continue

                if has_kl_column and len(parts) >= 4:
                    kl_str = parts[-3].replace(' ', '').replace('*', '')
                    if kl_str.replace('.', '').isdigit():
                        remaining = ' '.join(parts[:-3])
                    else:
                        remaining = ' '.join(parts[:-2])
                else:
                    remaining = ' '.join(parts[:-2])

                code, description = AudatexParser._split_code_and_description(
                    remaining)

                if description and rp_str:
                    result.append({
                        'code': code, 'description': description,
                        'rp': rp_str, 'cost': cost_str
                    })
            except (ValueError, IndexError):
                continue

        return result

    @staticmethod
    def _parse_paint(lines: List[str]) -> List[Dict[str, str]]:
        result = []
        section = AudatexParser._extract_section(
            lines,
            ['О К Р А С К А', 'ОКРАСКА'],
            ['ЛАКОКРАСОЧНЫЙ МАТЕРИАЛ', 'ЗАТРАТЫ ВРЕМЕНИ', 'ОКОНЧАТЕЛЬНАЯ']
        )

        for line in section:
            stripped = line.strip()
            if not stripped or stripped.startswith('-'):
                continue

            if any(h in stripped for h in [
                'КОД ОПЕР', 'СНЯТ ДЕТ', '2-СЛОЙН', 'СИСТЕМА AZT', 'СИСТЕМА AUDATEX',
                'ИТОГО', 'ВСЕГО', 'ЗАТРАТЫ ВРЕМЕНИ', 'СТОИМОСТЬ РАБОТ',
                'ВАЛЮТНЫЙ КУРС', 'НОРМА ВРЕМЕНИ', 'RUR/ЧАС', 'РЕМОНТ-КАЛЬКУЛЯЦИЯ', 'ЛИСТ',
                'ВРЕМЯ ОКРАСКИ', 'ПОДГОТОВКА', 'ПОДГ', 'ОБЩЕЕ ВРЕМЯ',
                'ОКРАСКА-ЗАТРАТЫ', 'МАТ.-КОНСТ', 'МАТЕРИАЛ-ИНДЕКС',
            ]):
                continue

            if re.search(r'\b\d+\s+(RUR|EUR|USD)\b', stripped):
                continue

            if AudatexParser._is_continuation_line(line):
                result.append({
                    'code': '', 'description': AudatexParser._normalize_spaces(stripped),
                    'rp': '', 'cost': ''
                })
                continue

            parts = line.split()
            if len(parts) < 2:
                continue

            rp_str = parts[-1].replace(' ', '').replace('*', '')
            if not rp_str.replace('.', '').isdigit():
                continue

            remaining = ' '.join(parts[:-1])
            code, description = AudatexParser._split_code_and_description(
                remaining)

            if description and rp_str:
                result.append({
                    'code': code, 'description': description,
                    'rp': rp_str, 'cost': ''
                })

        return result

    @staticmethod
    def _parse_materials(lines: List[str]) -> List[Dict[str, str]]:
        result = []
        section = AudatexParser._extract_section(
            lines,
            ['ЛАКОКРАСОЧНЫЙ МАТЕРИАЛ ЗА ДЕТАЛЬ', 'ЛАКОКРАСОЧНЫЙ МАТЕРИАЛ'],
            ['ЗАТРАТЫ ВРЕМЕНИ', 'ОКРАСКА-ЗАТРАТЫ', 'ОКОНЧАТЕЛЬНАЯ']
        )

        for line in section:
            stripped = line.strip()
            if not stripped or stripped.startswith('-'):
                continue

            if any(h in stripped for h in [
                'УПР №', 'НАЗВАНИЕ', 'СТОИМ.МАТ', 'СИСТЕМА AUDATEX',
                'ИТОГО', 'ВСЕГО', 'ЗАТРАТЫ ВРЕМЕНИ', 'СТОИМОСТЬ РАБОТ',
                'ВАЛЮТНЫЙ КУРС', 'RUR/ЧАС', 'ЛИСТ', 'РЕМОНТ-КАЛЬКУЛЯЦИЯ'
            ]):
                continue

            if AudatexParser._is_continuation_line(line):
                continue

            parts = line.split()
            if len(parts) < 2:
                continue

            try:
                cost_str = parts[-1].replace(' ', '').replace('*', '')
                if not cost_str.replace('.', '').isdigit():
                    continue

                remaining = ' '.join(parts[:-1])
                code, description = AudatexParser._split_code_and_description(
                    remaining)

                if description and cost_str:
                    result.append({
                        'code': code, 'description': description,
                        'rp': '', 'cost': cost_str
                    })
            except (ValueError, IndexError):
                continue

        return result


# ============================================================
# ДИАЛОГ ДОП. СТРОК ОКРАСКИ
# ============================================================

class ExtraPaintDialog(QDialog):
    def __init__(self, parent, paint_works: List[Dict], paint_materials: List[Dict]):
        super().__init__(parent)
        self.setWindowTitle("Дополнительные строки окраски")
        self.resize(700, 500)
        self.paint_works = paint_works
        self.paint_materials = paint_materials
        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        title = QLabel(
            "<b>Найдены дополнительные строки в блоках окраски.</b><br>"
            "Снимите галочки с ненужных строк, отредактируйте значения при необходимости."
        )
        title.setWordWrap(True)
        layout.addWidget(title)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)

        content = QWidget()
        content_layout = QVBoxLayout(content)
        content_layout.setSpacing(15)

        works_group = QGroupBox(
            f"✅ Работы по окраске (РП) — найдено {len(self.paint_works)}")
        works_layout = QVBoxLayout(works_group)
        self.works_checks = []
        self.works_values = []

        if self.paint_works:
            works_btns = QHBoxLayout()
            btn_all = QPushButton("Выбрать все")
            btn_none = QPushButton("Снять все")
            btn_all.clicked.connect(
                lambda: self._toggle_all(self.works_checks, True))
            btn_none.clicked.connect(
                lambda: self._toggle_all(self.works_checks, False))
            works_btns.addWidget(btn_all)
            works_btns.addWidget(btn_none)
            works_btns.addStretch()
            works_layout.addLayout(works_btns)

            for item in self.paint_works:
                row = QHBoxLayout()
                cb = QCheckBox()
                cb.setChecked(True)
                desc = QLineEdit(item['description'])
                desc.setMinimumWidth(300)
                val = QLineEdit(item['rp'])
                val.setFixedWidth(80)
                val.setAlignment(Qt.AlignmentFlag.AlignRight)

                row.addWidget(cb)
                row.addWidget(desc, 1)
                row.addWidget(QLabel("РП:"))
                row.addWidget(val)
                works_layout.addLayout(row)
                self.works_checks.append(cb)
                self.works_values.append((desc, val))
        else:
            works_layout.addWidget(QLabel("Строки не найдены"))

        content_layout.addWidget(works_group)

        mats_group = QGroupBox(
            f"✅ Материалы окраски (₽) — найдено {len(self.paint_materials)}")
        mats_layout = QVBoxLayout(mats_group)
        self.mats_checks = []
        self.mats_values = []

        if self.paint_materials:
            mats_btns = QHBoxLayout()
            btn_all = QPushButton("Выбрать все")
            btn_none = QPushButton("Снять все")
            btn_all.clicked.connect(
                lambda: self._toggle_all(self.mats_checks, True))
            btn_none.clicked.connect(
                lambda: self._toggle_all(self.mats_checks, False))
            mats_btns.addWidget(btn_all)
            mats_btns.addWidget(btn_none)
            mats_btns.addStretch()
            mats_layout.addLayout(mats_btns)

            for item in self.paint_materials:
                row = QHBoxLayout()
                cb = QCheckBox()
                cb.setChecked(True)
                desc = QLineEdit(item['description'])
                desc.setMinimumWidth(300)
                val = QLineEdit(item['cost'])
                val.setFixedWidth(100)
                val.setAlignment(Qt.AlignmentFlag.AlignRight)

                row.addWidget(cb)
                row.addWidget(desc, 1)
                row.addWidget(QLabel("₽:"))
                row.addWidget(val)
                mats_layout.addLayout(row)
                self.mats_checks.append(cb)
                self.mats_values.append((desc, val))
        else:
            mats_layout.addWidget(QLabel("Строки не найдены"))

        content_layout.addWidget(mats_group)
        content_layout.addStretch()

        scroll.setWidget(content)
        layout.addWidget(scroll, 1)

        btns = QHBoxLayout()
        btns.addStretch()
        btn_cancel = QPushButton("Отмена")
        btn_cancel.clicked.connect(self.reject)
        btn_ok = QPushButton("Добавить выбранные")
        btn_ok.setDefault(True)
        btn_ok.clicked.connect(self.accept)
        btns.addWidget(btn_cancel)
        btns.addWidget(btn_ok)
        layout.addLayout(btns)

    def _toggle_all(self, checks: List[QCheckBox], state: bool):
        for cb in checks:
            cb.setChecked(state)

    def get_selected(self) -> Dict[str, List[Dict[str, str]]]:
        result = {'works': [], 'materials': []}

        for cb, (desc_edit, val_edit) in zip(self.works_checks, self.works_values):
            if cb.isChecked():
                desc = desc_edit.text().strip()
                val = val_edit.text().strip()
                if desc and val:
                    try:
                        int(val)
                        result['works'].append(
                            {'description': desc, 'rp': val})
                    except ValueError:
                        pass

        for cb, (desc_edit, val_edit) in zip(self.mats_checks, self.mats_values):
            if cb.isChecked():
                desc = desc_edit.text().strip()
                val = val_edit.text().strip()
                if desc and val:
                    try:
                        int(val)
                        result['materials'].append(
                            {'description': desc, 'cost': val})
                    except ValueError:
                        pass

        return result


# ============================================================
# КОНСТАНТЫ
# ============================================================

BTN_PRIMARY = ("#3182ce", "white")
BTN_SUCCESS = ("#38a169", "white")
BTN_WARNING = ("#dd6b20", "white")
BTN_DANGER = ("#e53e3e", "white")
BTN_SAVE = ("#00796B", "white")
BTN_LOAD = ("#5D4037", "white")
BTN_LIGHT = ("#edf2f7", "#2d3748")
BTN_CALC = ("#9f7aea", "white")
BTN_MERGE = ("#ed64a6", "white")
BTN_PARSE = ("#d69e2e", "white")
BTN_EXTRA = ("#f59e0b", "white")

PASTE_COLORS = [
    "#c6f6d5", "#fefcbf", "#fed7d7", "#e9d8fd", "#bee3f8",
    "#feebc8", "#f687b3", "#68d391", "#fc8181", "#b794f4"
]

CLR_STATUS_OK = "#38a169"
CLR_STATUS_INFO = "#3182ce"
CLR_STATUS_TOTAL = "#38a169"
CLR_STATUS_AVG = "#805ad5"
CLR_STATUS_GRAND = "#2b6cb0"

CLR_HIGHLIGHT_BG = "#fff5f5"
CLR_HIGHLIGHT_BG_SEL = "#fed7d7"
CLR_HIGHLIGHT_FG = "#c53030"
CLR_SELECTION_BG = "#3182ce"

FONT_MAIN = ("Arial", 10)
FONT_MONO = ("Consolas", 11)
FONT_EXPORT = "Times New Roman"
FONT_EXPORT_SIZE = 12

CALC_VERSION = __version__

NUMERIC_COLUMNS = {
    "РП", "Стоимость", "Кол-во", "Износ", "Цена с износом",
    "Средняя цена", "Источник 1", "Источник 2", "Источник 3",
    "Источник 4", "Источник 5"
}

COLUMN_WIDTHS = {
    "Код": 100, "Артикул": 150, "РП": 80, "Стоимость": 100,
    "Кол-во": 80, "Износ": 80, "Цена с износом": 120,
    "Наименование": 250, "Средняя цена": 120, "Источник": 110
}

LIGHT_QSS = """
QMainWindow { background-color: #f8f9fa; }
QWidget#CentralWidget { background-color: #f8f9fa; }
QWidget { background-color: transparent; color: #212529; }
QDialog, QMessageBox, QInputDialog, QDialogButtonBox { background-color: #ffffff; color: #212529; }
QDialog QLabel, QMessageBox QLabel, QInputDialog QLabel { color: #212529; background: transparent; }
QMenu { background-color: #ffffff; border: 1px solid #dee2e6; border-radius: 6px; padding: 4px 0px; }
QMenu::item { background-color: transparent; color: #212529; padding: 6px 24px; }
QMenu::item:selected { background-color: #0d6efd; color: white; }
QMenu::separator { height: 1px; background: #dee2e6; margin: 4px 8px; }
QToolTip { background-color: #212529; color: #ffffff; border: 1px solid #212529; padding: 4px 8px; border-radius: 4px; font-size: 11px; }
QComboBox QAbstractItemView { background-color: #ffffff; color: #212529; selection-background-color: #0d6efd; selection-color: white; border: 1px solid #ced4da; outline: none; }
QTreeView, QListView, QTreeWidget, QListWidget { background-color: #ffffff; alternate-background-color: #f8f9fa; color: #212529; border: 1px solid #dee2e6; border-radius: 6px; }
QTreeView::item, QListView::item { padding: 4px; }
QTreeView::item:selected, QListView::item:selected { background-color: #0d6efd; color: white; }
QTabWidget::pane { border: 1px solid #dee2e6; background: white; border-radius: 6px; top: -1px; }
QTabBar::tab { background: #e9ecef; color: #495057; padding: 8px 16px; margin-right: 2px; border: 1px solid #dee2e6; border-bottom: none; border-top-left-radius: 6px; border-top-right-radius: 6px; }
QTabBar::tab:selected { background: white; color: #0d6efd; border-bottom: 2px solid white; margin-bottom: -1px; }
QTabBar::tab:hover:!selected { background: #f8f9fa; }
QTableWidget { background-color: white; alternate-background-color: #f8f9fa; gridline-color: #e9ecef; border: 1px solid #dee2e6; border-radius: 6px; selection-background-color: #e7f1ff; selection-color: #0d6efd; }
QTableWidget::item { padding: 6px; border-bottom: 1px solid #f1f3f5; }
QHeaderView::section { background-color: #f8f9fa; color: #212529; padding: 8px; border: none; border-bottom: 2px solid #dee2e6; border-right: 1px solid #e9ecef; font-weight: 600; }
QPushButton { border: 1px solid #ced4da; border-radius: 6px; padding: 6px 12px; background-color: white; color: #212529; font-weight: 500; min-height: 20px; }
QPushButton:hover { background-color: #e9ecef; border-color: #adb5bd; }
QPushButton:pressed { background-color: #dee2e6; }
QLineEdit, QComboBox { border: 1px solid #ced4da; border-radius: 6px; padding: 6px 8px; background: white; min-height: 20px; }
QLineEdit:focus, QComboBox:focus { border: 1px solid #86b7fe; outline: 0; }
QComboBox::drop-down { border: none; width: 20px; }
QGroupBox { border: 1px solid #dee2e6; border-radius: 6px; margin-top: 14px; padding-top: 14px; font-weight: 600; color: #495057; background: white; }
QGroupBox::title { subcontrol-origin: margin; subcontrol-position: top left; left: 12px; padding: 0 6px; color: #0d6efd; }
QLabel { color: #212529; }
QCheckBox { color: #212529; spacing: 6px; }
QCheckBox::indicator { width: 16px; height: 16px; border-radius: 4px; border: 1px solid #adb5bd; }
QCheckBox::indicator:checked { background-color: #0d6efd; border-color: #0d6efd; }
QScrollBar:vertical, QScrollBar:horizontal { background: #f8f9fa; width: 12px; height: 12px; margin: 0px; border-radius: 6px; border: 1px solid #e9ecef; }
QScrollBar::handle:vertical, QScrollBar::handle:horizontal { background: #ced4da; min-height: 20px; border-radius: 6px; }
QScrollBar::handle:vertical:hover, QScrollBar::handle:horizontal:hover { background: #adb5bd; }
QScrollBar::add-line, QScrollBar::sub-line { height: 0px; width: 0px; }
QScrollBar::add-page, QScrollBar::sub-page { background: none; }
QPlainTextEdit { background: white; border: 1px solid #ced4da; border-radius: 6px; padding: 4px; font-family: Consolas, monospace; }
"""

DARK_QSS = """
QMainWindow { background-color: #1a202c; color: #e2e8f0; }
QWidget#CentralWidget { background-color: #1a202c; color: #e2e8f0; }
QWidget { background-color: transparent; color: #e2e8f0; }
QDialog, QMessageBox, QInputDialog, QDialogButtonBox { background-color: #2d3748; color: #e2e8f0; }
QDialog QLabel, QMessageBox QLabel, QInputDialog QLabel { color: #e2e8f0; background: transparent; }
QMenu { background-color: #2d3748; border: 1px solid #4a5568; border-radius: 6px; padding: 4px 0px; }
QMenu::item { background-color: transparent; color: #e2e8f0; padding: 6px 24px; }
QMenu::item:selected { background-color: #4299e1; color: white; }
QMenu::separator { height: 1px; background: #4a5568; margin: 4px 8px; }
QToolTip { background-color: #1a202c; color: #e2e8f0; border: 1px solid #4a5568; padding: 4px 8px; border-radius: 4px; font-size: 11px; }
QComboBox QAbstractItemView { background-color: #2d3748; color: #e2e8f0; selection-background-color: #4299e1; selection-color: white; border: 1px solid #4a5568; outline: none; }
QTreeView, QListView, QTreeWidget, QListWidget { background-color: #2d3748; alternate-background-color: #252d3d; color: #e2e8f0; border: 1px solid #4a5568; border-radius: 6px; }
QTreeView::item, QListView::item { padding: 4px; }
QTreeView::item:selected, QListView::item:selected { background-color: #4299e1; color: white; }
QTabWidget::pane { border: 1px solid #2d3748; background: #2d3748; border-radius: 6px; top: -1px; }
QTabBar::tab { background: #1a202c; color: #a0aec0; padding: 8px 16px; margin-right: 2px; border: 1px solid #2d3748; border-bottom: none; border-top-left-radius: 6px; border-top-right-radius: 6px; }
QTabBar::tab:selected { background: #2d3748; color: #63b3ed; border-bottom: 2px solid #63b3ed; margin-bottom: -1px; }
QTabBar::tab:hover:!selected { background: #2d3748; }
QTableWidget { background-color: #2d3748; alternate-background-color: #252d3d; gridline-color: #4a5568; border: 1px solid #4a5568; border-radius: 6px; selection-background-color: #2b6cb0; selection-color: white; color: #e2e8f0; }
QTableWidget::item { padding: 6px; border-bottom: 1px solid #2a2a2a; }
QHeaderView::section { background-color: #1a202c; color: #cbd5e0; padding: 8px; border: none; border-bottom: 2px solid #4a5568; border-right: 1px solid #2d3748; font-weight: 600; }
QPushButton { border: 1px solid #4a5568; border-radius: 6px; padding: 6px 12px; background-color: #2d3748; color: #e2e8f0; font-weight: 500; min-height: 20px; }
QPushButton:hover { background-color: #4a5568; border-color: #718096; }
QPushButton:pressed { background-color: #1a202c; }
QLineEdit, QComboBox { border: 1px solid #4a5568; border-radius: 6px; padding: 6px 8px; background: #1a202c; color: #e2e8f0; min-height: 20px; }
QLineEdit:focus, QComboBox:focus { border: 1px solid #4299e1; }
QComboBox::drop-down { border: none; width: 20px; }
QGroupBox { border: 1px solid #4a5568; border-radius: 6px; margin-top: 14px; padding-top: 14px; font-weight: 600; color: #cbd5e0; background: #2d3748; }
QGroupBox::title { subcontrol-origin: margin; subcontrol-position: top left; left: 12px; padding: 0 6px; color: #63b3ed; }
QLabel { color: #e2e8f0; }
QCheckBox { color: #e2e8f0; spacing: 6px; }
QCheckBox::indicator { width: 16px; height: 16px; border-radius: 4px; border: 1px solid #555555; background: #1a202c; }
QCheckBox::indicator:checked { background-color: #4299e1; border-color: #4299e1; }
QScrollBar:vertical, QScrollBar:horizontal { background: #1a202c; width: 12px; height: 12px; margin: 0px; border-radius: 6px; }
QScrollBar::handle:vertical, QScrollBar::handle:horizontal { background: #4a5568; min-height: 20px; border-radius: 6px; }
QScrollBar::handle:vertical:hover, QScrollBar::handle:horizontal:hover { background: #718096; }
QScrollBar::add-line, QScrollBar::sub-line { height: 0px; width: 0px; }
QScrollBar::add-page, QScrollBar::sub-page { background: none; }
QPlainTextEdit { background: #1a202c; border: 1px solid #4a5568; border-radius: 6px; padding: 4px; font-family: Consolas, monospace; color: #e2e8f0; }
"""


@contextmanager
def blocked_signals(*widgets):
    for w in widgets:
        w.blockSignals(True)
    try:
        yield
    finally:
        for w in widgets:
            w.blockSignals(False)


@contextmanager
def wait_cursor():
    QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
    try:
        yield
    finally:
        QApplication.restoreOverrideCursor()


def is_numeric_column(header: str) -> bool:
    return header in NUMERIC_COLUMNS


def parse_price(text: str) -> float:
    if not text:
        return 0.0
    cleaned = re.sub(r"[^\d.,\-]", "", text.strip()).replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def create_styled_button(text: str, bg_color: str, text_color: str, bold: bool = True, padding: int = 6) -> QPushButton:
    btn = QPushButton(text)
    font_weight = "600" if bold else "normal"
    btn.setStyleSheet(f"""
        QPushButton {{ background-color: {bg_color}; color: {text_color}; font-weight: {font_weight}; padding: {padding}px 14px; border: none; border-radius: 6px; min-height: 20px; }}
        QPushButton:hover {{ background-color: {bg_color}dd; }}
        QPushButton:pressed {{ background-color: {bg_color}aa; }}
    """)
    return btn


def get_column_width(header: str) -> int:
    if header.startswith("Источник"):
        return COLUMN_WIDTHS["Источник"]
    return COLUMN_WIDTHS.get(header, 100)


class SplashScreen(QSplashScreen):
    def __init__(self, image_path: str = None):
        pixmap = self._load_or_create_pixmap(image_path)
        super().__init__(pixmap)
        self.setWindowFlag(Qt.WindowType.WindowStaysOnTopHint)

        self.loading_label = QLabel("Загрузка...", self)
        pixmap_width = pixmap.width()
        self.loading_label.setGeometry(
            50, pixmap.height() - 65, pixmap_width - 100, 20)
        self.loading_label.setStyleSheet(
            "color: #333333; font-size: 11px; font-weight: bold; background: transparent;")
        self.loading_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.progress_bar = QProgressBar(self)
        self.progress_bar.setGeometry(
            50, pixmap.height() - 40, pixmap_width - 100, 25)
        self.progress_bar.setStyleSheet(
            "QProgressBar { border: 2px solid #444; border-radius: 3px; text-align: center; background-color: rgba(255,255,255,0.5); color: white; font-weight: bold; } QProgressBar::chunk { background-color: #00BCD4; border-radius: 2px; }")
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat("")

        self.version_label = QLabel(f"v{CALC_VERSION}", self)
        self.version_label.setGeometry(
            pixmap_width - 100, pixmap.height() - 15, 100, 15)
        self.version_label.setStyleSheet(
            "color: #00BCD4; font-size: 9px; font-weight: bold; background: transparent;")
        self.version_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.author_label = QLabel(f"\u00a9 {__author__}", self)
        self.author_label.setGeometry(50, pixmap.height() - 15, 300, 15)
        self.author_label.setStyleSheet(
            "color: #666; font-size: 8px; background: transparent;")
        self.author_label.setAlignment(Qt.AlignmentFlag.AlignLeft)

    def _load_or_create_pixmap(self, image_path: str) -> QPixmap:
        if image_path and os.path.exists(image_path):
            try:
                pixmap = QPixmap(image_path)
                if not pixmap.isNull():
                    return pixmap.scaled(pixmap.width() // 2, pixmap.height() // 2, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
            except Exception:
                pass
        return self._create_fallback_pixmap()

    def _create_fallback_pixmap(self) -> QPixmap:
        pixmap = QPixmap(400, 210)
        pixmap.fill(QColor("#F5F5F5"))
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.fillRect(pixmap.rect(), QColor("#E3F2FD"))
        painter.setPen(QColor("#90CAF9"))
        painter.setBrush(QColor("#BBDEFB"))
        painter.drawRoundedRect(50, 75, 300, 100, 10, 10)
        painter.drawRoundedRect(100, 50, 200, 50, 8, 8)
        painter.setBrush(QColor("#424242"))
        painter.drawEllipse(90, 150, 30, 30)
        painter.drawEllipse(280, 150, 30, 30)
        painter.setBrush(QColor("#90CAF9"))
        painter.drawEllipse(90, 150, 15, 15)
        painter.drawEllipse(280, 150, 15, 15)
        painter.setBrush(QColor("#E3F2FD"))
        painter.drawRoundedRect(110, 55, 75, 35, 5, 5)
        painter.drawRoundedRect(215, 55, 75, 35, 5, 5)
        painter.setPen(QColor("#003366"))
        font = QFont("Arial", 18, QFont.Weight.Bold)
        painter.setFont(font)
        painter.drawText(QRect(50, 20, 300, 30),
                         Qt.AlignmentFlag.AlignCenter, "Audatex")
        font = QFont("Arial", 12, QFont.Weight.Bold)
        painter.setFont(font)
        painter.setPen(QColor("#0277BD"))
        painter.drawText(QRect(50, 45, 300, 20),
                         Qt.AlignmentFlag.AlignCenter, "Converter")
        painter.setPen(QColor("#666666"))
        font = QFont("Arial", 8)
        painter.setFont(font)
        painter.drawText(QRect(50, 190, 200, 15), Qt.AlignmentFlag.AlignLeft |
                         Qt.AlignmentFlag.AlignVCenter, f"v{CALC_VERSION}")
        painter.drawText(QRect(250, 190, 125, 15), Qt.AlignmentFlag.AlignRight |
                         Qt.AlignmentFlag.AlignVCenter, f"\u00a9 {__author__}")
        painter.end()
        return pixmap

    def update_progress(self, value: int, message: str = ""):
        self.progress_bar.setValue(value)
        if message:
            self.loading_label.setText(message)


class ConfigManager:
    CONFIG_FILE = os.path.join(os.path.dirname(
        os.path.abspath(__file__)), "config.json")

    def __init__(self):
        self.config = {}
        self.load()

    def load(self):
        if os.path.exists(self.CONFIG_FILE):
            try:
                with open(self.CONFIG_FILE, "r", encoding="utf-8") as f:
                    self.config = json.load(f)
            except Exception:
                self.config = {}

    def save(self):
        try:
            with open(self.CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(self.config, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def get_dict_path(self) -> str:
        return self.config.get("dict_path", "")

    def set_dict_path(self, path: str):
        self.config["dict_path"] = path
        self.save()


class AbbrevManager:
    def __init__(self, dict_path: str = ""):
        self.dict_path = dict_path
        self.abbreviations: Dict[str, str] = {}
        self.load()

    def set_dict_path(self, path: str):
        self.dict_path = path
        self.load()

    def load(self):
        if self.dict_path and os.path.exists(self.dict_path):
            try:
                with open(self.dict_path, "r", encoding="utf-8") as f:
                    self.abbreviations = json.load(f)
            except Exception:
                self.abbreviations = {}
        else:
            self.abbreviations = {}

    def save(self) -> bool:
        if not self.dict_path:
            return False
        try:
            with open(self.dict_path, "w", encoding="utf-8") as f:
                json.dump(self.abbreviations, f, ensure_ascii=False, indent=2)
            return True
        except Exception:
            return False

    def add(self, abbr: str, exp: str):
        self.abbreviations[abbr.upper().strip()] = exp.strip()
        self.save()

    def remove(self, abbr: str):
        key = abbr.upper().strip()
        if key in self.abbreviations:
            del self.abbreviations[key]
            self.save()

    def get_all(self) -> Dict[str, str]:
        return self.abbreviations.copy()

    def is_known(self, text: str) -> bool:
        return text.upper().strip() in self.abbreviations

    def has_upper_phrases(self, text: str) -> bool:
        pattern = r"(?<![А-ЯЁA-Za-zа-яё])(?:[А-ЯЁA-Z0-9]{2,}|[А-ЯЁA-Z0-9]+(?:[\s/\-]+[А-ЯЁA-Z0-9]+)+)(?![А-ЯЁA-Za-zа-яё])"
        return bool(re.search(pattern, text))

    def should_highlight(self, text: str) -> bool:
        if not text or not text.strip():
            return False
        if not self.has_upper_phrases(text):
            return False
        if self.is_known(text):
            return False
        return True

    def find_unknown(self, text: str) -> List[str]:
        pattern = r"(?<![А-ЯЁA-Za-zа-яё])(?:[А-ЯЁA-Z0-9]{2,}|[А-ЯЁA-Z0-9]+(?:[\s/\-]+[А-ЯЁA-Z0-9]+)+)(?![А-ЯЁA-Za-zа-яё])"
        unknown = []
        seen = set()
        for match in re.finditer(pattern, text):
            phrase = match.group().strip()
            if phrase not in seen and not self.is_known(phrase):
                unknown.append(phrase)
                seen.add(phrase)
        return unknown

    def expand(self, text: str) -> str:
        if not self.abbreviations:
            return text
        result = text
        for old in sorted(self.abbreviations.keys(), key=len, reverse=True):
            exp = self.abbreviations[old]
            pattern = r"(?<![А-ЯЁA-Za-zа-яё])" + \
                re.escape(old) + r"(?![А-ЯЁA-Za-zа-яё])"
            result = re.sub(pattern, exp, result, flags=re.IGNORECASE)
        return result


class AbbrevDelegate(QStyledItemDelegate):
    def __init__(self, parent=None, manager: AbbrevManager = None, highlight_column: str = "Описание", is_dark=False):
        super().__init__(parent)
        self.manager = manager
        self.highlight_column = highlight_column
        self.is_dark = is_dark

    def set_dark_mode(self, is_dark: bool):
        self.is_dark = is_dark

    def paint(self, painter: QPainter, option: QStyleOptionViewItem, index: QModelIndex):
        self.initStyleOption(option, index)
        if index.model().headerData(index.column(), Qt.Orientation.Horizontal) != self.highlight_column:
            super().paint(painter, option, index)
            return

        text = index.data()
        if not text or not self.manager or not self.manager.should_highlight(str(text)):
            super().paint(painter, option, index)
            return

        painter.save()
        if self.is_dark:
            bg = "#5c2c2c" if option.state & QStyle.StateFlag.State_Selected else "#4a2020"
            fg = "#ff9999"
        else:
            bg = CLR_HIGHLIGHT_BG_SEL if option.state & QStyle.StateFlag.State_Selected else CLR_HIGHLIGHT_BG
            fg = CLR_HIGHLIGHT_FG

        painter.fillRect(option.rect, QColor(bg))
        painter.setPen(QColor(fg))
        font = option.font
        font.setBold(True)
        painter.setFont(font)
        painter.drawText(
            option.rect.adjusted(4, 2, -4, -2),
            Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter | Qt.TextFlag.TextWordWrap,
            str(text)
        )
        painter.restore()


class ColumnSelectPlainTextEdit(QPlainTextEdit):
    def __init__(self, parent=None):
        super().__init__(parent)
        self._rect_selecting = False
        self._rect_start = None
        self._last_mouse_pos = None
        self._selected_texts: List[str] = []
        self._scroll_timer = QTimer(self)
        self._scroll_timer.setInterval(30)
        self._scroll_timer.timeout.connect(self._auto_scroll_tick)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and event.modifiers() & Qt.KeyboardModifier.AltModifier:
            self._rect_selecting = True
            self._rect_start = self.cursorForPosition(event.pos())
            self._last_mouse_pos = event.pos()
            self._update_rect_selection()
            self._scroll_timer.start()
            event.accept()
        else:
            self._rect_selecting = False
            self._selected_texts = []
            self.setExtraSelections([])
            self._scroll_timer.stop()
            super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self._rect_selecting:
            self._last_mouse_pos = event.pos()
            self._update_rect_selection()
            event.accept()
        else:
            super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self._rect_selecting:
            self._rect_selecting = False
            self._scroll_timer.stop()
            event.accept()
        else:
            super().mouseReleaseEvent(event)

    def keyPressEvent(self, event):
        key = event.key()
        modifiers = event.modifiers()
        is_copy = ((key == Qt.Key.Key_C and modifiers & Qt.KeyboardModifier.ControlModifier) or
                   (key == Qt.Key.Key_Insert and modifiers & Qt.KeyboardModifier.ControlModifier))
        if is_copy and self._selected_texts:
            QApplication.clipboard().setText('\n'.join(self._selected_texts))
            event.accept()
            return
        super().keyPressEvent(event)

    def _auto_scroll_tick(self):
        if not self._rect_selecting or self._last_mouse_pos is None:
            return
        viewport = self.viewport()
        vbar = self.verticalScrollBar()
        hbar = self.horizontalScrollBar()
        margin, scroll_step = 30, 5
        rect = viewport.rect()
        y, x = self._last_mouse_pos.y(), self._last_mouse_pos.x()
        scrolled = False
        if y < margin and vbar.value() > vbar.minimum():
            vbar.setValue(vbar.value() - scroll_step)
            scrolled = True
        elif y > rect.height() - margin and vbar.value() < vbar.maximum():
            vbar.setValue(vbar.value() + scroll_step)
            scrolled = True
        if x < margin and hbar.value() > hbar.minimum():
            hbar.setValue(hbar.value() - scroll_step)
            scrolled = True
        elif x > rect.width() - margin and hbar.value() < hbar.maximum():
            hbar.setValue(hbar.value() + scroll_step)
            scrolled = True
        if scrolled:
            self._update_rect_selection()

    def _update_rect_selection(self):
        if not self._rect_start:
            return
        current_cursor = self.cursorForPosition(
            self._last_mouse_pos) if self._last_mouse_pos else self._rect_start
        start_block = self._rect_start.blockNumber()
        end_block = current_cursor.blockNumber()
        start_col = self._rect_start.position() - self._rect_start.block().position()
        end_col = current_cursor.position() - current_cursor.block().position()
        b_start, b_end = min(start_block, end_block), max(
            start_block, end_block)
        c_start, c_end = min(start_col, end_col), max(start_col, end_col)
        extra_selections = []
        self._selected_texts = []
        for b in range(b_start, b_end + 1):
            block = self.document().findBlockByNumber(b)
            if not block.isValid():
                self._selected_texts.append("")
                continue
            line_len = len(block.text())
            if c_start >= line_len:
                self._selected_texts.append("")
                continue
            actual_end = min(c_end, line_len)
            cursor = QTextCursor(block)
            cursor.setPosition(block.position() + c_start)
            length = actual_end - c_start
            if length > 0:
                cursor.movePosition(
                    QTextCursor.MoveOperation.Right, QTextCursor.MoveMode.KeepAnchor, length)
            sel = QTextEdit.ExtraSelection()
            sel.cursor = cursor
            sel.format.setBackground(QColor(CLR_SELECTION_BG))
            sel.format.setForeground(QColor("white"))
            extra_selections.append(sel)
            self._selected_texts.append(
                cursor.selectedText() if length > 0 else "")
        self.setExtraSelections(extra_selections)

    def copy(self):
        if self._selected_texts:
            QApplication.clipboard().setText('\n'.join(self._selected_texts))
        else:
            super().copy()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"AUDATEX → LibreOffice Converter v{CALC_VERSION}")
        self.resize(1200, 800)
        self.is_dark_mode = False
        self.apply_theme(False)

        self.config_mgr = ConfigManager()
        self._manager = None

        self.rate = 285.0
        self.rp_work = 10
        self.rp_paint = 10
        self.wear_percent = 0.0
        self.small_parts_percent = 0.0
        self.total_labels: Dict[QWidget, QLabel] = {}
        self.last_saved_path = None

        self._setup_ui()
        QTimer.singleShot(100, self._init_manager)

    def _init_manager(self):
        dict_path = self.config_mgr.get_dict_path()
        if dict_path and os.path.exists(dict_path):
            self._manager = AbbrevManager(dict_path)
            self.status.setText(
                f"✅ Словарь загружен: {os.path.basename(dict_path)}")
        else:
            self._manager = AbbrevManager("")
            self.status.setText(
                "⚠️ Словарь не найден. Нажмите «📖 Словарь» для выбора.")
        self._refresh_all_tables()

    @property
    def manager(self) -> Optional[AbbrevManager]:
        return self._manager

    def apply_theme(self, is_dark: bool):
        self.is_dark_mode = is_dark
        app = QApplication.instance()
        app.setStyleSheet(DARK_QSS if is_dark else LIGHT_QSS)
        palette = QPalette()
        if is_dark:
            palette.setColor(QPalette.ColorRole.Window, QColor("#1a202c"))
            palette.setColor(QPalette.ColorRole.WindowText, QColor("#e2e8f0"))
            palette.setColor(QPalette.ColorRole.Base, QColor("#2d3748"))
            palette.setColor(QPalette.ColorRole.AlternateBase,
                             QColor("#252d3d"))
            palette.setColor(QPalette.ColorRole.ToolTipBase, QColor("#1a202c"))
            palette.setColor(QPalette.ColorRole.ToolTipText, QColor("#e2e8f0"))
            palette.setColor(QPalette.ColorRole.Text, QColor("#e2e8f0"))
            palette.setColor(QPalette.ColorRole.Button, QColor("#2d3748"))
            palette.setColor(QPalette.ColorRole.ButtonText, QColor("#e2e8f0"))
            palette.setColor(QPalette.ColorRole.BrightText, QColor("#fc8181"))
            palette.setColor(QPalette.ColorRole.Link, QColor("#63b3ed"))
            palette.setColor(QPalette.ColorRole.Highlight, QColor("#4299e1"))
            palette.setColor(
                QPalette.ColorRole.HighlightedText, QColor("#ffffff"))
        else:
            palette.setColor(QPalette.ColorRole.Window, QColor("#f8f9fa"))
            palette.setColor(QPalette.ColorRole.WindowText, QColor("#212529"))
            palette.setColor(QPalette.ColorRole.Base, QColor("#ffffff"))
            palette.setColor(QPalette.ColorRole.AlternateBase,
                             QColor("#f8f9fa"))
            palette.setColor(QPalette.ColorRole.ToolTipBase, QColor("#ffffff"))
            palette.setColor(QPalette.ColorRole.ToolTipText, QColor("#212529"))
            palette.setColor(QPalette.ColorRole.Text, QColor("#212529"))
            palette.setColor(QPalette.ColorRole.Button, QColor("#ffffff"))
            palette.setColor(QPalette.ColorRole.ButtonText, QColor("#212529"))
            palette.setColor(QPalette.ColorRole.BrightText, QColor("#e53e3e"))
            palette.setColor(QPalette.ColorRole.Link, QColor("#0d6efd"))
            palette.setColor(QPalette.ColorRole.Highlight, QColor("#0d6efd"))
            palette.setColor(
                QPalette.ColorRole.HighlightedText, QColor("#ffffff"))
        app.setPalette(palette)

        for tab in [getattr(self, 'works', None), getattr(self, 'paint', None), getattr(self, 'mats', None), getattr(self, 'parts', None)]:
            if tab:
                t = tab.findChild(QTableWidget)
                if t and t.itemDelegate():
                    t.itemDelegate().set_dark_mode(is_dark)

        if hasattr(self, 'total_summary'):
            bg = "#2b6cb0" if is_dark else CLR_STATUS_GRAND
            self.total_summary.setStyleSheet(
                f"background-color: {bg}; color: white; font-weight: 600; font-size: 14px; padding: 8px; border-radius: 6px;")
        if hasattr(self, 'avg_total_label'):
            bg = "#553c9a" if is_dark else CLR_STATUS_AVG
            self.avg_total_label.setStyleSheet(
                f"background-color: {bg}; color: white; font-weight: 600; font-size: 12px; padding: 6px; border-radius: 6px;")
        if hasattr(self, 'status'):
            bg = "#2d3748" if is_dark else "#e9ecef"
            clr = "#68d391" if is_dark else CLR_STATUS_OK
            self.status.setStyleSheet(
                f"color:{clr}; font-weight:600; padding:8px; background: {bg}; border-radius: 6px; margin-top: 5px;")

    def _setup_ui(self):
        central = QWidget()
        central.setObjectName("CentralWidget")
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)
        layout.setContentsMargins(15, 15, 15, 15)
        layout.setSpacing(10)
        self._build_top_panel(layout)
        self._build_dict_path_panel(layout)
        self._build_params_panel(layout)
        self._build_hint(layout)
        self._build_tabs(layout)
        self._build_status_bar(layout)

    def _build_top_panel(self, parent_layout):
        top = QHBoxLayout()
        btn_html = create_styled_button("🌐 Загрузить HTML", *BTN_PRIMARY)
        btn_html.clicked.connect(self.load_html)
        top.addWidget(btn_html)
        btn_dict = create_styled_button("📖 Словарь", *BTN_WARNING)
        btn_dict.clicked.connect(self.open_dict_editor)
        top.addWidget(btn_dict)
        self.chk_expand = QCheckBox("Расшифровывать при экспорте")
        self.chk_expand.setChecked(True)
        top.addWidget(self.chk_expand)
        self.chk_dark_mode = QCheckBox("🌙 Тёмная тема")
        self.chk_dark_mode.stateChanged.connect(
            lambda state: self.apply_theme(state == Qt.CheckState.Checked.value))
        top.addWidget(self.chk_dark_mode)
        top.addSpacing(20)
        btn_save = create_styled_button("💾 Сохранить", *BTN_SAVE)
        btn_save.clicked.connect(self._save_calculation)
        top.addWidget(btn_save)
        btn_load = create_styled_button("📂 Загрузить", *BTN_LOAD)
        btn_load.clicked.connect(self._load_calculation)
        top.addWidget(btn_load)
        top.addStretch()
        self.btn_export = create_styled_button("💾 Экспорт ▼", *BTN_SUCCESS)
        self.btn_export.setMenu(self._create_export_menu())
        top.addWidget(self.btn_export)
        parent_layout.addLayout(top)

    def _build_dict_path_panel(self, parent_layout):
        db_layout = QHBoxLayout()
        self.lbl_dict_path = QLabel()
        self._update_dict_path_label()
        self.lbl_dict_path.setStyleSheet(
            f"color:{CLR_STATUS_INFO}; font-weight:600; padding:4px;")
        db_layout.addWidget(self.lbl_dict_path)
        btn_change_db = create_styled_button(
            " Изменить путь", *BTN_LIGHT, padding=4)
        btn_change_db.clicked.connect(self._change_dict_path)
        db_layout.addWidget(btn_change_db)
        db_layout.addStretch()
        parent_layout.addLayout(db_layout)

    def _build_params_panel(self, parent_layout):
        params = QGroupBox("⚙️ Расчёт стоимости")
        grid = QGridLayout(params)
        grid.setHorizontalSpacing(15)
        grid.setVerticalSpacing(10)
        grid.addWidget(QLabel("Стоимость нормо-часа (₽):"), 0, 0)
        self.in_rate = QLineEdit("285")
        self.in_rate.setFixedWidth(100)
        self.in_rate.textChanged.connect(self._update_rate)
        grid.addWidget(self.in_rate, 0, 1)
        grid.addWidget(QLabel("РП/ч (работы):"), 0, 2)
        self.in_rp_w = QComboBox()
        self.in_rp_w.addItems(["10", "12"])
        self.in_rp_w.setMinimumWidth(80)
        self.in_rp_w.currentTextChanged.connect(self._update_rp_work)
        grid.addWidget(self.in_rp_w, 0, 3)
        grid.addWidget(QLabel("РП/ч (окраска):"), 0, 4)
        self.in_rp_p = QComboBox()
        self.in_rp_p.addItems(["10", "12"])
        self.in_rp_p.setMinimumWidth(80)
        self.in_rp_p.currentTextChanged.connect(self._update_rp_paint)
        grid.addWidget(self.in_rp_p, 0, 5)
        grid.addWidget(QLabel("Износ (%):"), 1, 0)
        self.in_wear = QLineEdit("0")
        self.in_wear.setFixedWidth(100)
        self.in_wear.setPlaceholderText("0-100")
        self.in_wear.textChanged.connect(self._update_wear)
        grid.addWidget(self.in_wear, 1, 1)
        grid.addWidget(QLabel("Мелкие детали (%):"), 1, 2)
        self.in_small_parts = QLineEdit("0")
        self.in_small_parts.setFixedWidth(100)
        self.in_small_parts.setPlaceholderText("0-100")
        self.in_small_parts.textChanged.connect(
            self._update_small_parts_percent)
        grid.addWidget(self.in_small_parts, 1, 3)
        btn_calc = create_styled_button("🔄 Пересчитать всё", *BTN_CALC)
        btn_calc.clicked.connect(self._recalc_all)
        grid.addWidget(btn_calc, 1, 4, 1, 2)
        parent_layout.addWidget(params)

    def _build_hint(self, parent_layout):
        hint = QLabel(
            "💡 🔴 Красным подсвечиваются слова/фразы ВЕРХНЕГО РЕГИСТРА, которых нет в словаре. "
            "Alt+мышь = выделение столбцом. Ctrl+C / Ctrl+Ins — копирование. "
            "ПКМ в запчастях → «🔄 Перенести в материалы». ПКМ в материалах → «🔄 Вернуть в запчасти»."
        )
        hint.setStyleSheet(
            "background:#edf2f7; padding:10px; border-radius:6px; color:#2d3748; border: 1px solid #e2e8f0;")
        hint.setWordWrap(True)
        parent_layout.addWidget(hint)

    def _build_tabs(self, parent_layout):
        self.tabs = QTabWidget()
        parent_layout.addWidget(self.tabs, 1)
        raw_tab = QWidget()
        raw_layout = QVBoxLayout(raw_tab)
        parse_btns = QHBoxLayout()
        btn_parse = create_styled_button(
            "🔍 Автоматически распарсить текст", *BTN_PARSE, padding=8)
        btn_parse.clicked.connect(self.auto_parse_text)
        parse_btns.addWidget(btn_parse)
        btn_extra = create_styled_button(
            "🔍 Найти доп. окраски", *BTN_EXTRA, padding=8)
        btn_extra.clicked.connect(self._find_extra_paint)
        parse_btns.addWidget(btn_extra)
        parse_btns.addStretch()
        raw_layout.addLayout(parse_btns)
        self.editor = ColumnSelectPlainTextEdit()
        self.editor.setFont(QFont(*FONT_MONO))
        self.editor.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        raw_layout.addWidget(self.editor)
        self.tabs.addTab(raw_tab, " Сырой текст")
        self.works = self._create_table(
            "Работы", ["Код", "Описание", "РП", "Стоимость"], True, "Описание", False)
        self.tabs.addTab(self.works, "🔧 Работы")
        self.paint = self._create_table(
            "Окраска", ["Код", "Описание", "РП", "Стоимость"], True, "Описание", False)
        self.tabs.addTab(self.paint, "🎨 Окраска-работы")
        self.mats = self._create_table(
            "Материалы", ["Код", "Описание", "Стоимость"], False, "Описание", False)
        self.tabs.addTab(self.mats, "🧪 Окраска-материалы")
        self.parts = self._create_table(
            "Запчасти",
            ["Код", "Кол-во", "Наименование", "Артикул",
                "Стоимость", "Износ", "Цена с износом"],
            False, "Наименование", True
        )
        self.tabs.addTab(self.parts, "🔩 Запчасти")
        self.avg_prices_tab = self._create_avg_prices_tab()
        self.tabs.addTab(self.avg_prices_tab, "📊 Средние цены")
        self.total_summary = QLabel("💰 Общий итог: 0 ₽")
        self.total_summary.setAlignment(Qt.AlignmentFlag.AlignCenter)
        parent_layout.addWidget(self.total_summary)
        self.apply_theme(self.is_dark_mode)

    def _build_status_bar(self, parent_layout):
        self.status = QLabel("✅ Готово.")
        self.status.setAlignment(Qt.AlignmentFlag.AlignCenter)
        parent_layout.addWidget(self.status)
        self.apply_theme(self.is_dark_mode)

    def auto_parse_text(self):
        text = self.editor.toPlainText()
        if not text.strip():
            QMessageBox.warning(self, "Внимание", "Текст пуст!")
            return
        with wait_cursor():
            try:
                parsed = AudatexParser.parse(text)
                works_table = self.works.findChild(QTableWidget)
                if works_table:
                    with blocked_signals(works_table):
                        works_table.setRowCount(0)
                        for item in parsed['works']:
                            row = works_table.rowCount()
                            works_table.insertRow(row)
                            works_table.setItem(
                                row, 0, QTableWidgetItem(item['code']))
                            works_table.setItem(
                                row, 1, QTableWidgetItem(item['description']))
                            works_table.setItem(
                                row, 2, QTableWidgetItem(item['rp']))
                            works_table.setItem(
                                row, 3, QTableWidgetItem(item['cost']))
                paint_table = self.paint.findChild(QTableWidget)
                if paint_table:
                    with blocked_signals(paint_table):
                        paint_table.setRowCount(0)
                        for item in parsed['paint']:
                            row = paint_table.rowCount()
                            paint_table.insertRow(row)
                            paint_table.setItem(
                                row, 0, QTableWidgetItem(item['code']))
                            paint_table.setItem(
                                row, 1, QTableWidgetItem(item['description']))
                            paint_table.setItem(
                                row, 2, QTableWidgetItem(item['rp']))
                            paint_table.setItem(
                                row, 3, QTableWidgetItem(item['cost']))
                mats_table = self.mats.findChild(QTableWidget)
                if mats_table:
                    with blocked_signals(mats_table):
                        mats_table.setRowCount(0)
                        for item in parsed['materials']:
                            row = mats_table.rowCount()
                            mats_table.insertRow(row)
                            mats_table.setItem(
                                row, 0, QTableWidgetItem(item['code']))
                            mats_table.setItem(
                                row, 1, QTableWidgetItem(item['description']))
                            mats_table.setItem(
                                row, 2, QTableWidgetItem(item['cost']))
                parts_table = self.parts.findChild(QTableWidget)
                if parts_table:
                    with blocked_signals(parts_table):
                        parts_table.setRowCount(0)
                        for item in parsed['parts']:
                            row = parts_table.rowCount()
                            parts_table.insertRow(row)
                            parts_table.setItem(
                                row, 0, QTableWidgetItem(item['code']))
                            parts_table.setItem(
                                row, 1, QTableWidgetItem(item['qty']))
                            parts_table.setItem(
                                row, 2, QTableWidgetItem(item['name']))
                            parts_table.setItem(
                                row, 3, QTableWidgetItem(item['article']))
                            parts_table.setItem(
                                row, 4, QTableWidgetItem(item['cost']))
                self._recalc_all()
                for table in [works_table, paint_table, mats_table, parts_table]:
                    if table:
                        table.viewport().update()
                msg = (
                    f"✅ Текст успешно распарсен!\n"
                    f"• Запчастей: {len(parsed['parts'])}\n"
                    f"• Работ: {len(parsed['works'])}\n"
                    f"• Окраска: {len(parsed['paint'])}\n"
                    f"• Материалов: {len(parsed['materials'])}\n"
                    f"💡 Совет: нажмите «🔍 Найти доп. окраски» для добавления "
                    f"итоговых строк (ПОДГ, МАТ.-КОНСТ. и др.)"
                )
                QMessageBox.information(self, "Успех", msg)
                self.status.setText(
                    f"✅ Распарсено: {len(parsed['parts'])} запч., {len(parsed['works'])} работ, "
                    f"{len(parsed['paint'])} окр., {len(parsed['materials'])} мат."
                )
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка парсинга", f"Не удалось распарсить текст:\n{e}")

    def _find_extra_paint(self):
        text = self.editor.toPlainText()
        if not text.strip():
            QMessageBox.warning(
                self, "Внимание", "Сначала загрузите или вставьте текст!")
            return
        with wait_cursor():
            extra = AudatexParser.find_extra_paint_lines(text)
            works_count = len(extra['works'])
            mats_count = len(extra['materials'])
            if works_count == 0 and mats_count == 0:
                QMessageBox.information(
                    self, "Ничего не найдено",
                    "В тексте не найдены блоки «ЗАТРАТЫ ВРЕМЕНИ НА ОКРАСКУ» "
                    "или «ОКРАСКА-ЗАТРАТЫ НА МАТЕРИАЛ».\n"
                    "Возможно, в этой смете они отсутствуют или имеют другой формат."
                )
                return
            dialog = ExtraPaintDialog(self, extra['works'], extra['materials'])
            if dialog.exec() != QDialog.DialogCode.Accepted:
                return
            selected = dialog.get_selected()
            added_works = len(selected['works'])
            added_mats = len(selected['materials'])
            if added_works == 0 and added_mats == 0:
                QMessageBox.information(
                    self, "Ничего не добавлено", "Вы не выбрали ни одной строки.")
                return
            self._add_extra_paint_lines(selected)
            msg = (
                f"✅ Добавлено дополнительных строк:\n"
                f"• Работ по окраске: {added_works}\n"
                f"• Материалов: {added_mats}"
            )
            QMessageBox.information(self, "Успех", msg)
            self.status.setText(
                f"✅ Добавлено: {added_works} работ окр., {added_mats} мат.")

    def _add_extra_paint_lines(self, selected: Dict[str, List[Dict[str, str]]]):
        paint_table = self.paint.findChild(QTableWidget)
        mats_table = self.mats.findChild(QTableWidget)
        if selected['works'] and paint_table:
            with blocked_signals(paint_table):
                for item in selected['works']:
                    row = paint_table.rowCount()
                    paint_table.insertRow(row)
                    paint_table.setItem(row, 0, QTableWidgetItem(""))
                    paint_table.setItem(
                        row, 1, QTableWidgetItem(item['description']))
                    paint_table.setItem(row, 2, QTableWidgetItem(item['rp']))
                    paint_table.setItem(row, 3, QTableWidgetItem(""))
                paint_table.viewport().update()
        if selected['materials'] and mats_table:
            with blocked_signals(mats_table):
                for item in selected['materials']:
                    row = mats_table.rowCount()
                    mats_table.insertRow(row)
                    mats_table.setItem(row, 0, QTableWidgetItem(""))
                    mats_table.setItem(
                        row, 1, QTableWidgetItem(item['description']))
                    mats_table.setItem(row, 2, QTableWidgetItem(item['cost']))
                mats_table.viewport().update()
        self._recalc_all()

    def _clear_table(self, tab):
        table = tab.findChild(QTableWidget)
        if not table or table.rowCount() == 0:
            return
        if QMessageBox.question(self, "Очистить?", "Удалить все строки из таблицы?") == QMessageBox.StandardButton.Yes:
            with blocked_signals(table):
                table.setRowCount(0)
            self._update_table_total(tab)
            self._update_total_summary()

    def _create_avg_prices_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(5, 5, 5, 5)
        sources_group = QGroupBox("🌐 Источники информации (URL сайтов)")
        sources_layout = QVBoxLayout(sources_group)
        self.source_edits: List[QLineEdit] = []
        for i in range(5):
            h = QHBoxLayout()
            h.addWidget(QLabel(f"Источник {i+1}:"))
            edit = QLineEdit()
            edit.setPlaceholderText(f"https://example{i+1}.com")
            self.source_edits.append(edit)
            h.addWidget(edit, 1)
            sources_layout.addLayout(h)
        layout.addWidget(sources_group)
        btns = QHBoxLayout()
        btn_add = create_styled_button(
            "➕ Добавить строку", *BTN_LIGHT, padding=4)
        btn_add.clicked.connect(lambda: self.avg_prices_table.insertRow(
            self.avg_prices_table.rowCount()) if hasattr(self, 'avg_prices_table') else None)
        btns.addWidget(btn_add)
        btn_del = create_styled_button(
            "➖ Удалить строку", *BTN_LIGHT, padding=4)
        btn_del.clicked.connect(self._delete_avg_price_row)
        btns.addWidget(btn_del)
        btn_clear = create_styled_button("🗑️ Очистить", *BTN_LIGHT, padding=4)
        btn_clear.clicked.connect(self._clear_avg_prices)
        btns.addWidget(btn_clear)
        btn_export_avg = create_styled_button(
            "📤 Экспорт средних цен ▼", *BTN_SUCCESS, padding=4)
        btn_export_avg.setMenu(self._create_avg_export_menu())
        btns.addWidget(btn_export_avg)
        btns.addStretch()
        btn_apply = create_styled_button(
            "📥 Использовать средние цены", *BTN_SUCCESS, padding=8)
        btn_apply.clicked.connect(self._apply_avg_prices)
        btns.addWidget(btn_apply)
        layout.addLayout(btns)
        headers = ["Наименование", "Артикул", "Источник 1", "Источник 2",
                   "Источник 3", "Источник 4", "Источник 5", "Средняя цена"]
        table = QTableWidget(0, len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.setShowGrid(True)
        table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows)
        table.itemChanged.connect(self._on_avg_price_item_changed)
        table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Interactive)
        table.horizontalHeader().setMinimumSectionSize(50)
        for i, h in enumerate(headers):
            if h == "Наименование":
                table.horizontalHeader().setSectionResizeMode(i, QHeaderView.ResizeMode.Stretch)
            else:
                table.setColumnWidth(i, get_column_width(h))
        layout.addWidget(table, 1)
        self.avg_prices_table = table
        self.avg_total_label = QLabel("💰 Средняя сумма: 0 ₽")
        self.avg_total_label.setAlignment(
            Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        layout.addWidget(self.avg_total_label)
        return tab

    def _create_avg_export_menu(self) -> QMenu:
        menu = QMenu(self)
        for label, handler in [
            ("📊 LibreOffice Calc (.ods)", self._export_avg_ods),
            ("📝 LibreOffice Writer (.odt)", self._export_avg_odt),
            ("📈 Microsoft Excel (.xlsx)", self._export_avg_excel),
            ("📄 Microsoft Word (.docx)", self._export_avg_word),
        ]:
            act = QAction(label, self)
            act.triggered.connect(handler)
            menu.addAction(act)
        return menu

    def _get_avg_sources_info(self) -> List[str]:
        sources = []
        if hasattr(self, 'source_edits'):
            for i, edit in enumerate(self.source_edits):
                url = edit.text().strip()
                if url:
                    sources.append(f"Источник {i+1} — {url}")
        return sources

    def _get_avg_table_data(self) -> Tuple[List[str], List[List[str]]]:
        headers = ["Наименование", "Артикул", "Источник 1", "Источник 2",
                   "Источник 3", "Источник 4", "Источник 5", "Средняя цена"]
        data = []
        if hasattr(self, 'avg_prices_table'):
            for row_idx in range(self.avg_prices_table.rowCount()):
                row = []
                for col_idx in range(len(headers)):
                    item = self.avg_prices_table.item(row_idx, col_idx)
                    row.append(item.text() if item else "")
                data.append(row)
        return headers, data

    def _export_avg_excel(self):
        if not hasattr(self, 'avg_prices_table') or self.avg_prices_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Таблица средних цен пуста!")
            return
        if not HAS_XLSX:
            QMessageBox.critical(
                self, "Ошибка", "Установите: pip install openpyxl")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Экспорт средних цен", "avg_prices.xlsx", "Excel файлы (*.xlsx);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Средние цены"
                headers, data = self._get_avg_table_data()
                header_font = XlFont(
                    name=FONT_EXPORT, size=FONT_EXPORT_SIZE, bold=True, color="FFFFFF")
                header_fill = PatternFill(
                    start_color="2196F3", end_color="2196F3", fill_type="solid")
                border = Border(left=Side(style='thin'), right=Side(
                    style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.border = border
                    cell.alignment = Alignment(
                        horizontal='center', vertical='center')
                for row_idx, row in enumerate(data):
                    for col_idx, value in enumerate(row):
                        cell = ws.cell(
                            row=row_idx+2, column=col_idx+1, value=value)
                        cell.border = border
                        if col_idx >= 2:
                            cell.alignment = Alignment(
                                horizontal='right', vertical='center')
                            try:
                                cell.value = int(value) if value else ""
                            except (ValueError, TypeError):
                                pass
                        else:
                            cell.alignment = Alignment(
                                horizontal='left', vertical='center')
                sources = self._get_avg_sources_info()
                if sources:
                    start_row = len(data) + 4
                    title_cell = ws.cell(
                        row=start_row, column=1, value="Источники информации:")
                    title_cell.font = XlFont(
                        name=FONT_EXPORT, size=FONT_EXPORT_SIZE, bold=True, color="003366")
                    for i, source in enumerate(sources):
                        source_cell = ws.cell(
                            row=start_row + 1 + i, column=1, value=source)
                        source_cell.font = XlFont(
                            name=FONT_EXPORT, size=FONT_EXPORT_SIZE)
                        source_cell.alignment = Alignment(
                            horizontal='left', vertical='center')
                for col in ws.columns:
                    max_length = max((len(str(cell.value))
                                     for cell in col if cell.value), default=0)
                    ws.column_dimensions[col[0].column_letter].width = min(
                        max_length + 2, 50)
                wb.save(path)
                self.status.setText(
                    f"✅ Экспортированы средние цены: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Средние цены экспортированы:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось экспортировать:\n{e}")

    def _export_avg_ods(self):
        if not hasattr(self, 'avg_prices_table') or self.avg_prices_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Таблица средних цен пуста!")
            return
        if not HAS_ODF:
            QMessageBox.critical(
                self, "Ошибка", "Установите: pip install odfpy")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Экспорт средних цен", "avg_prices.ods", "ODS файлы (*.ods);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = OpenDocumentSpreadsheet()
                text_style = Style(name="TS_TNR12", family="table-cell")
                text_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt"))
                text_style.addElement(ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(text_style)
                num_style = Style(name="TS_TNR12_Right", family="table-cell")
                num_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt"))
                num_style.addElement(ParagraphProperties(
                    lineheight="100%", textalign="end"))
                doc.automaticstyles.addElement(num_style)
                bold_style = Style(name="TS_TNR12_Bold", family="table-cell")
                bold_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt", fontweight="bold"))
                bold_style.addElement(ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(bold_style)
                headers, data = self._get_avg_table_data()
                tbl = Table(name="Средние цены")
                for _ in headers:
                    tbl.addElement(TableColumn())
                r = TableRow()
                for h in headers:
                    c = TableCell(valuetype="string",
                                  stylename="TS_TNR12_Bold")
                    c.addElement(P(text=h))
                    r.addElement(c)
                tbl.addElement(r)
                for row in data:
                    r = TableRow()
                    for i, val in enumerate(row):
                        style = "TS_TNR12_Right" if i >= 2 else "TS_TNR12"
                        c = TableCell(valuetype="string", stylename=style)
                        c.addElement(P(text=str(val)))
                        r.addElement(c)
                    tbl.addElement(r)
                doc.spreadsheet.addElement(tbl)
                sources = self._get_avg_sources_info()
                if sources:
                    src_tbl = Table(name="Источники")
                    src_tbl.addElement(TableColumn())
                    r = TableRow()
                    c = TableCell(valuetype="string",
                                  stylename="TS_TNR12_Bold")
                    c.addElement(P(text="Источники информации:"))
                    r.addElement(c)
                    src_tbl.addElement(r)
                    for source in sources:
                        r = TableRow()
                        c = TableCell(valuetype="string", stylename="TS_TNR12")
                        c.addElement(P(text=source))
                        r.addElement(c)
                        src_tbl.addElement(r)
                    doc.spreadsheet.addElement(src_tbl)
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортированы средние цены: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Средние цены экспортированы:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось экспортировать:\n{e}")

    def _export_avg_odt(self):
        if not hasattr(self, 'avg_prices_table') or self.avg_prices_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Таблица средних цен пуста!")
            return
        if not HAS_ODF:
            QMessageBox.critical(
                self, "Ошибка", "Установите: pip install odfpy")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Экспорт средних цен", "avg_prices.odt", "ODT файлы (*.odt);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = OpenDocumentText()
                title_style = Style(name="Title", family="paragraph")
                title_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize="18pt", fontweight="bold"))
                title_style.addElement(ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(title_style)
                heading_style = Style(
                    name="Heading_Custom", family="paragraph")
                heading_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize="14pt", fontweight="bold"))
                heading_style.addElement(
                    ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(heading_style)
                cell_style = Style(name="Cell_TNR12", family="table-cell")
                cell_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt"))
                cell_style.addElement(ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(cell_style)
                cell_right_style = Style(
                    name="Cell_TNR12_Right", family="table-cell")
                cell_right_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt"))
                cell_right_style.addElement(ParagraphProperties(
                    lineheight="100%", textalign="end"))
                doc.automaticstyles.addElement(cell_right_style)
                cell_bold_style = Style(
                    name="Cell_TNR12_Bold", family="table-cell")
                cell_bold_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt", fontweight="bold"))
                cell_bold_style.addElement(
                    ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(cell_bold_style)
                doc.text.addElement(
                    P(text="Средние цены на запчасти", stylename="Title"))
                headers, data = self._get_avg_table_data()
                tbl = Table(name="Средние цены")
                for _ in headers:
                    tbl.addElement(TableColumn())
                r = TableRow()
                for h in headers:
                    c = TableCell(valuetype="string",
                                  stylename="Cell_TNR12_Bold")
                    c.addElement(P(text=h))
                    r.addElement(c)
                tbl.addElement(r)
                for row in data:
                    r = TableRow()
                    for i, val in enumerate(row):
                        style = "Cell_TNR12_Right" if i >= 2 else "Cell_TNR12"
                        c = TableCell(valuetype="string", stylename=style)
                        c.addElement(P(text=str(val)))
                        r.addElement(c)
                    tbl.addElement(r)
                doc.text.addElement(tbl)
                sources = self._get_avg_sources_info()
                if sources:
                    doc.text.addElement(P(text=""))
                    doc.text.addElement(
                        P(text="Источники информации:", stylename="Heading_Custom"))
                    for source in sources:
                        doc.text.addElement(P(text=source))
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортированы средние цены: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Средние цены экспортированы:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось экспортировать:\n{e}")

    def _export_avg_word(self):
        if not hasattr(self, 'avg_prices_table') or self.avg_prices_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Таблица средних цен пуста!")
            return
        if not HAS_DOCX:
            QMessageBox.critical(
                self, "Ошибка", "Установите: pip install python-docx")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Экспорт средних цен", "avg_prices.docx", "Word файлы (*.docx);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = DocxDocument()
                title = doc.add_heading('Средние цены на запчасти', level=0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._apply_docx_font(title, size=18, bold=True)
                headers, data = self._get_avg_table_data()
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                for i, h in enumerate(headers):
                    hdr_cells[i].text = h
                    for paragraph in hdr_cells[i].paragraphs:
                        self._apply_docx_paragraph_format(paragraph)
                        for run in paragraph.runs:
                            self._apply_docx_run_font(run, bold=True)
                for row in data:
                    row_cells = table.add_row().cells
                    for i, val in enumerate(row):
                        row_cells[i].text = str(val)
                        for paragraph in row_cells[i].paragraphs:
                            self._apply_docx_paragraph_format(paragraph)
                            for run in paragraph.runs:
                                self._apply_docx_run_font(run, bold=False)
                sources = self._get_avg_sources_info()
                if sources:
                    doc.add_paragraph()
                    heading = doc.add_heading('Источники информации:', level=1)
                    self._apply_docx_font(heading, size=14, bold=True)
                    for source in sources:
                        p = doc.add_paragraph(source)
                        self._apply_docx_paragraph_format(p)
                        for run in p.runs:
                            self._apply_docx_run_font(run, bold=False)
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортированы средние цены: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Средние цены экспортированы:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось экспортировать:\n{e}")

    def _create_table(self, name, headers, calc_cost, highlight_column, is_parts=False):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(5, 5, 5, 5)
        btns = QHBoxLayout()
        btn_add = create_styled_button("➕ Добавить", *BTN_LIGHT, padding=4)
        btn_add.clicked.connect(lambda: self._on_table_changed(tab))
        btns.addWidget(btn_add)
        btn_del = create_styled_button("➖ Удалить", *BTN_LIGHT, padding=4)
        btn_del.clicked.connect(lambda: self._delete_selected_rows(tab))
        btns.addWidget(btn_del)
        btn_clear = create_styled_button("🗑️ Очистить", *BTN_LIGHT, padding=4)
        btn_clear.clicked.connect(lambda checked, t=tab: self._clear_table(t))
        btns.addWidget(btn_clear)
        btn_merge = create_styled_button(
            "🔗 Объединить продолжения", *BTN_MERGE, padding=4)
        btn_merge.clicked.connect(
            lambda: self._merge_continuations(tab.findChild(QTableWidget)))
        btns.addWidget(btn_merge)
        btns.addStretch()
        layout.addLayout(btns)
        paste = QGroupBox("📋 Вставить в колонку:")
        paste_l = QHBoxLayout(paste)
        for i, h in enumerate(headers):
            if is_parts and h == "Износ":
                continue
            if is_parts and h == "Цена с износом":
                b = create_styled_button(
                    "📥 Перенести в средние цены", *BTN_DANGER, padding=4)
                b.clicked.connect(self._transfer_parts_to_avg_prices)
            else:
                bg = PASTE_COLORS[i % len(PASTE_COLORS)]
                b = create_styled_button(f"📋 В «{h}»", bg, "black", padding=4)
                b.clicked.connect(
                    lambda checked, h=h: self._paste_to_col(tab, h))
            paste_l.addWidget(b)
        layout.addWidget(paste)
        table = QTableWidget(0, len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.setShowGrid(True)
        table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows)
        table.itemChanged.connect(
            lambda item: self._on_item_changed(item, calc_cost, tab))
        table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Interactive)
        table.horizontalHeader().setMinimumSectionSize(50)
        for i, h in enumerate(headers):
            if h in ("Описание", "Наименование"):
                table.horizontalHeader().setSectionResizeMode(i, QHeaderView.ResizeMode.Stretch)
            else:
                table.setColumnWidth(i, get_column_width(h))
        table.setItemDelegate(AbbrevDelegate(
            table, self._manager, highlight_column, self.is_dark_mode))
        table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        table.customContextMenuRequested.connect(
            lambda pos: self._show_table_context_menu(table, pos, highlight_column))
        layout.addWidget(table)
        total_label = QLabel("💰 Итого: 0 ₽")
        total_label.setStyleSheet(
            f"background-color: {CLR_STATUS_TOTAL}; color: white; font-weight: bold; "
            f"font-size: 12px; padding: 6px; border-radius: 4px;"
        )
        total_label.setAlignment(
            Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        layout.addWidget(total_label)
        self.total_labels[tab] = total_label
        table.calc_cost = calc_cost
        table.name = name
        table.highlight_column = highlight_column
        table.is_parts = is_parts
        return tab

    def _create_export_menu(self) -> QMenu:
        menu = QMenu(self)
        for label, handler in [
            ("📊 LibreOffice Calc (.ods)", self.export_ods),
            ("📝 LibreOffice Writer (.odt)", self.export_odt),
            ("📈 Microsoft Excel (.xlsx)", self.export_excel),
            ("📄 Microsoft Word (.docx)", self.export_word),
        ]:
            act = QAction(label, self)
            act.triggered.connect(handler)
            menu.addAction(act)
        return menu

    def _on_table_changed(self, tab):
        self._update_table_total(tab)
        self._update_total_summary()

    def _update_table_total(self, tab):
        if tab not in self.total_labels:
            return
        table = tab.findChild(QTableWidget)
        if not table:
            return
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        if getattr(table, 'is_parts', False) and "Цена с износом" in headers:
            cost_col_name = "Цена с износом"
        elif "Стоимость" in headers:
            cost_col_name = "Стоимость"
        else:
            self.total_labels[tab].setText("💰 Итого: —")
            return
        cost_idx = headers.index(cost_col_name)
        total = count = 0
        for r in range(table.rowCount()):
            item = table.item(r, cost_idx)
            if item:
                try:
                    total += int(item.text().strip())
                    count += 1
                except (ValueError, TypeError):
                    pass
        total_str = f"{total:,}".replace(",", " ")
        self.total_labels[tab].setText(
            f"💰 Итого: {total_str} ₽   ({count} поз.)")

    def _update_total_summary(self):
        grand_total = grand_count = 0
        for tab in [self.works, self.paint, self.mats, self.parts]:
            table = tab.findChild(QTableWidget)
            if not table:
                continue
            headers = [table.horizontalHeaderItem(
                i).text() for i in range(table.columnCount())]
            if getattr(table, 'is_parts', False) and "Цена с износом" in headers:
                cost_col_name = "Цена с износом"
            elif "Стоимость" in headers:
                cost_col_name = "Стоимость"
            else:
                continue
            cost_idx = headers.index(cost_col_name)
            for r in range(table.rowCount()):
                item = table.item(r, cost_idx)
                if item:
                    try:
                        grand_total += int(item.text().strip())
                        grand_count += 1
                    except (ValueError, TypeError):
                        pass
        total_str = f"{grand_total:,}".replace(",", " ")
        self.total_summary.setText(
            f"💰 Общий итог: {total_str} ₽   ({grand_count} позиций)")

    def _delete_selected_rows(self, tab):
        table = tab.findChild(QTableWidget)
        rows = set(idx.row() for idx in table.selectedIndexes())
        for row in sorted(rows, reverse=True):
            table.removeRow(row)
        self._update_table_total(tab)
        self._update_total_summary()

    def _delete_avg_price_row(self):
        if not hasattr(self, 'avg_prices_table'):
            return
        rows = set(idx.row()
                   for idx in self.avg_prices_table.selectedIndexes())
        for row in sorted(rows, reverse=True):
            self.avg_prices_table.removeRow(row)
        self._update_avg_total()

    def _clear_avg_prices(self):
        if not hasattr(self, 'avg_prices_table') or self.avg_prices_table.rowCount() == 0:
            return
        if QMessageBox.question(self, "Очистить?", "Удалить все строки средних цен?") == QMessageBox.StandardButton.Yes:
            self.avg_prices_table.setRowCount(0)
            self._update_avg_total()

    def _show_table_context_menu(self, table: QTableWidget, pos: QPoint, highlight_column: str):
        item = table.itemAt(pos)
        if not item:
            return
        col_header = table.horizontalHeaderItem(item.column()).text()
        menu = QMenu(self)
        is_parts_table = getattr(table, 'is_parts', False)
        is_mats_table = (table.name == "Материалы")
        if col_header == highlight_column and self.manager and self.manager.should_highlight(item.text()):
            act = QAction(f"📝 Внести «{item.text()}» в словарь", self)
            act.triggered.connect(
                lambda checked: self._quick_add(item.text(), table))
            menu.addAction(act)
            menu.addSeparator()
        act_copy_unknown = QAction("📋 Копировать в буфер нераспознанные", self)
        act_copy_unknown.triggered.connect(
            lambda checked: self._copy_unknown_to_clipboard(table, highlight_column))
        menu.addAction(act_copy_unknown)
        menu.addSeparator()
        if is_parts_table:
            act_to_mats = QAction("🔄 Перенести в материалы", self)
            act_to_mats.triggered.connect(
                lambda checked: self._move_parts_to_materials(table))
            menu.addAction(act_to_mats)
            menu.addSeparator()
        if is_mats_table:
            act_to_parts = QAction("🔄 Вернуть в запчасти", self)
            act_to_parts.triggered.connect(
                lambda checked: self._move_materials_to_parts(table))
            menu.addAction(act_to_parts)
            menu.addSeparator()
        for label, handler in [
            ("Копировать", self._copy_selection),
            ("Вырезать", self._cut_selection),
            ("Вставить", self._paste_selection),
        ]:
            act = QAction(label, self)
            act.triggered.connect(lambda checked, h=handler: h(table))
            menu.addAction(act)
        menu.exec(table.mapToGlobal(pos))

    def _move_parts_to_materials(self, parts_table: QTableWidget):
        selected_rows = sorted(
            set(idx.row() for idx in parts_table.selectedIndexes()), reverse=True)
        if not selected_rows:
            QMessageBox.information(
                self, "Внимание", "Сначала выделите строки для переноса!")
            return
        mats_table = self.mats.findChild(QTableWidget)
        if not mats_table:
            return
        parts_headers = [parts_table.horizontalHeaderItem(
            i).text() for i in range(parts_table.columnCount())]
        try:
            code_idx = parts_headers.index("Код")
            qty_idx = parts_headers.index("Кол-во")
            name_idx = parts_headers.index("Наименование")
            article_idx = parts_headers.index("Артикул")
            cost_idx = parts_headers.index("Стоимость")
        except ValueError:
            QMessageBox.critical(
                self, "Ошибка", "Не удалось найти нужные колонки!")
            return
        moved = 0
        with blocked_signals(parts_table, mats_table):
            for row in selected_rows:
                code_item = parts_table.item(row, code_idx)
                qty_item = parts_table.item(row, qty_idx)
                name_item = parts_table.item(row, name_idx)
                article_item = parts_table.item(row, article_idx)
                cost_item = parts_table.item(row, cost_idx)
                code = code_item.text().strip() if code_item else ""
                qty = qty_item.text().strip() if qty_item else "1"
                name = name_item.text().strip() if name_item else ""
                article = article_item.text().strip() if article_item else ""
                cost = cost_item.text().strip() if cost_item else ""
                if not name:
                    continue
                parts = []
                if article:
                    parts.append(f"({article})")
                if qty and qty != "1":
                    parts.append(f"- {qty} шт.")
                if parts:
                    description = f"{name} {' '.join(parts)}"
                else:
                    description = name
                new_row = mats_table.rowCount()
                mats_table.insertRow(new_row)
                mats_table.setItem(new_row, 0, QTableWidgetItem(code))
                mats_table.setItem(new_row, 1, QTableWidgetItem(description))
                mats_table.setItem(new_row, 2, QTableWidgetItem(cost))
                parts_table.removeRow(row)
                moved += 1
            parts_table.viewport().update()
            mats_table.viewport().update()
        self._update_table_total(self.parts)
        self._update_table_total(self.mats)
        self._update_total_summary()
        QMessageBox.information(
            self, "Готово", f"Перенесено строк из запчастей в материалы: {moved}")

    def _move_materials_to_parts(self, mats_table: QTableWidget):
        selected_rows = sorted(
            set(idx.row() for idx in mats_table.selectedIndexes()), reverse=True)
        if not selected_rows:
            QMessageBox.information(
                self, "Внимание", "Сначала выделите строки для переноса!")
            return
        parts_table = self.parts.findChild(QTableWidget)
        if not parts_table:
            return
        mats_headers = [mats_table.horizontalHeaderItem(
            i).text() for i in range(mats_table.columnCount())]
        try:
            code_idx = mats_headers.index("Код")
            desc_idx = mats_headers.index("Описание")
            cost_idx = mats_headers.index("Стоимость")
        except ValueError:
            QMessageBox.critical(
                self, "Ошибка", "Не удалось найти нужные колонки!")
            return
        parts_headers = [parts_table.horizontalHeaderItem(
            i).text() for i in range(parts_table.columnCount())]
        try:
            p_code_idx = parts_headers.index("Код")
            p_qty_idx = parts_headers.index("Кол-во")
            p_name_idx = parts_headers.index("Наименование")
            p_article_idx = parts_headers.index("Артикул")
            p_cost_idx = parts_headers.index("Стоимость")
        except ValueError:
            QMessageBox.critical(
                self, "Ошибка", "Не удалось найти нужные колонки в запчастях!")
            return
        moved = 0
        with blocked_signals(parts_table, mats_table):
            for row in selected_rows:
                code_item = mats_table.item(row, code_idx)
                desc_item = mats_table.item(row, desc_idx)
                cost_item = mats_table.item(row, cost_idx)
                code = code_item.text().strip() if code_item else ""
                description = desc_item.text().strip() if desc_item else ""
                cost = cost_item.text().strip() if cost_item else ""
                if not description:
                    continue
                article = ""
                qty = "1"
                name = description
                qty_match = re.search(
                    r'\s*-\s*(\d+)\s*шт\.?\s*$', name, re.IGNORECASE)
                if qty_match:
                    qty = qty_match.group(1)
                    name = name[:qty_match.start()].strip()
                article_match = re.search(r'\s*\(([^)]+)\)\s*$', name)
                if article_match:
                    article = article_match.group(1).strip()
                    name = name[:article_match.start()].strip()
                new_row = parts_table.rowCount()
                parts_table.insertRow(new_row)
                parts_table.setItem(new_row, p_code_idx,
                                    QTableWidgetItem(code))
                parts_table.setItem(new_row, p_qty_idx, QTableWidgetItem(qty))
                parts_table.setItem(new_row, p_name_idx,
                                    QTableWidgetItem(name))
                parts_table.setItem(new_row, p_article_idx,
                                    QTableWidgetItem(article))
                parts_table.setItem(new_row, p_cost_idx,
                                    QTableWidgetItem(cost))
                mats_table.removeRow(row)
                moved += 1
            parts_table.viewport().update()
            mats_table.viewport().update()
        self._apply_wear_to_all_parts()
        self._update_table_total(self.parts)
        self._update_table_total(self.mats)
        self._update_total_summary()
        QMessageBox.information(
            self, "Готово", f"Возвращено строк из материалов в запчасти: {moved}")

    def _copy_unknown_to_clipboard(self, table: QTableWidget, highlight_column: str):
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        try:
            col_idx = headers.index(highlight_column)
        except ValueError:
            return
        unknown_set = set()
        for r in range(table.rowCount()):
            item = table.item(r, col_idx)
            if not item:
                continue
            text = item.text().strip()
            if text and self.manager and self.manager.should_highlight(text):
                unknown_set.add(text)
        if not unknown_set:
            QMessageBox.information(
                self, "Информация", "Все фразы в таблице распознаны! 🎉")
            return
        QApplication.clipboard().setText("\n".join(sorted(unknown_set)))
        QMessageBox.information(
            self, "✅ Скопировано в буфер",
            f"Скопировано уникальных нераспознанных фраз: {len(unknown_set)}\n"
            f"Теперь можно вставить в ИИ (Ctrl+V) для расшифровки, "
            f"а затем добавить в словарь через «📖 Словарь»."
        )

    def _copy_selection(self, table):
        items = table.selectedItems()
        if items:
            QApplication.clipboard().setText("\n".join(i.text() for i in items))

    def _cut_selection(self, table):
        self._copy_selection(table)
        for i in table.selectedItems():
            i.setText("")

    def _paste_selection(self, table):
        text = QApplication.clipboard().text()
        if text and table.selectedItems():
            table.selectedItems()[0].setText(text)

    def _quick_add(self, phrase: str, table: QTableWidget):
        if not self.manager or not self.manager.dict_path:
            QMessageBox.warning(self, "Внимание", "База данных не выбрана!")
            return
        exp, ok = QInputDialog.getText(
            self, "Добавить сокращение", f"Расшифровка для: {phrase}")
        if ok and exp:
            self.manager.add(phrase, exp)
            table.viewport().update()
            QMessageBox.information(self, "✅", f"Добавлено: {phrase}")

    def _paste_to_col(self, tab, col_name):
        text = QApplication.clipboard().text()
        if not text:
            return QMessageBox.warning(self, "Внимание", "Буфер пуст")
        lines = text.split("\n")
        while lines and lines[-1].strip() == "":
            lines.pop()
        if not lines:
            return QMessageBox.warning(self, "Внимание", "Нет данных")
        table = tab.findChild(QTableWidget)
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        try:
            col_idx = headers.index(col_name)
        except ValueError:
            return QMessageBox.critical(self, "Ошибка", f"Колонка '{col_name}' не найдена")
        has_data = any(
            table.item(r, col_idx) and table.item(r, col_idx).text().strip()
            for r in range(table.rowCount())
        )
        start_row = 0
        if has_data:
            reply = QMessageBox.question(
                self, "Колонка заполнена",
                f"В «{col_name}» уже есть данные.\n"
                f"• Да — продолжить внесение (добавить после)\n"
                f"• Нет — внести заново (очистить колонку)\n"
                f"• Отмена — не вставлять",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Cancel
            )
            if reply == QMessageBox.StandardButton.Cancel:
                return
            elif reply == QMessageBox.StandardButton.Yes:
                start_row = table.rowCount()
                for r in range(table.rowCount()):
                    if not table.item(r, col_idx) or not table.item(r, col_idx).text().strip():
                        start_row = r
                        break
            else:
                with blocked_signals(table):
                    for r in range(table.rowCount()):
                        if table.item(r, col_idx):
                            table.item(r, col_idx).setText("")
        with blocked_signals(table):
            inserted = updated = 0
            for i, line in enumerate(lines):
                clean_value = re.sub(r"\s{2,}", " ", line).strip()
                row_idx = start_row + i
                if row_idx < table.rowCount():
                    item = table.item(row_idx, col_idx) or QTableWidgetItem()
                    table.setItem(row_idx, col_idx, item)
                    item.setText(clean_value)
                    updated += 1
                else:
                    table.insertRow(row_idx)
                    for c in range(table.columnCount()):
                        table.setItem(row_idx, c, QTableWidgetItem(
                            clean_value if c == col_idx else ""))
                    inserted += 1
            table.viewport().update()
            if getattr(table, 'is_parts', False):
                self._recalc_parts_table(table)
            elif table.calc_cost and col_name == "РП":
                self._recalc_table(table)
            self._update_table_total(tab)
            self._update_total_summary()
        QMessageBox.information(
            self, "Готово",
            f"Колонка: «{col_name}»\n"
            f"Обновлено: {updated}\n"
            f"Добавлено: {inserted}\n"
            f"Всего: {updated + inserted}"
        )

    def _on_item_changed(self, item, calc, tab):
        table = item.tableWidget()
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        col_header = headers[item.column()] if item.column() < len(
            headers) else ""
        if getattr(table, 'is_parts', False):
            if col_header in ("Стоимость", "Износ"):
                self._recalc_parts_row(table, item.row())
        elif calc and col_header == "РП":
            self._recalc_row(table, item.row())
        table.viewport().update()
        self._update_table_total(tab)
        self._update_total_summary()

    def _on_avg_price_item_changed(self, item):
        table = item.tableWidget()
        self._recalc_avg_price_row(table, item.row())
        self._update_avg_total()

    def _recalc_row(self, table, row):
        rp_item = table.item(row, 2)
        if not rp_item:
            return
        try:
            rp = int(rp_item.text().strip())
            rp_h = self.rp_paint if table.name == "Окраска" else self.rp_work
            cost = round((rp / rp_h) * self.rate)
            with blocked_signals(table):
                ci = table.item(row, 3) or QTableWidgetItem()
                table.setItem(row, 3, ci)
                ci.setText(str(cost))
        except (ValueError, TypeError):
            pass

    def _recalc_parts_row(self, table, row):
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        try:
            cost_idx = headers.index("Стоимость")
            wear_idx = headers.index("Износ")
            price_with_wear_idx = headers.index("Цена с износом")
        except ValueError:
            return
        cost_item = table.item(row, cost_idx)
        if not cost_item:
            return
        try:
            cost = float(cost_item.text().strip())
        except (ValueError, TypeError):
            return
        wear_item = table.item(row, wear_idx)
        if not wear_item or not wear_item.text().strip():
            with blocked_signals(table):
                if not wear_item:
                    wear_item = QTableWidgetItem()
                table.setItem(row, wear_idx, wear_item)
                wear_item.setText(str(self.wear_percent))
                wear = self.wear_percent
        else:
            try:
                wear = float(wear_item.text().strip())
            except (ValueError, TypeError):
                wear = self.wear_percent
        wear = max(0.0, min(100.0, wear))
        price_with_wear = round(cost * (1 - wear / 100))
        with blocked_signals(table):
            pww_item = table.item(
                row, price_with_wear_idx) or QTableWidgetItem()
            table.setItem(row, price_with_wear_idx, pww_item)
            pww_item.setText(str(price_with_wear))

    def _recalc_avg_price_row(self, table: QTableWidget, row: int):
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        if "Средняя цена" not in headers:
            return
        avg_idx = headers.index("Средняя цена")
        prices = []
        for i in range(1, 6):
            col_name = f"Источник {i}"
            if col_name in headers:
                col_idx = headers.index(col_name)
                item = table.item(row, col_idx)
                if item and item.text().strip():
                    price = parse_price(item.text())
                    if price > 0:
                        prices.append(price)
        avg_price = round(sum(prices) / len(prices)) if prices else 0
        with blocked_signals(table):
            avg_item = table.item(row, avg_idx) or QTableWidgetItem()
            table.setItem(row, avg_idx, avg_item)
            avg_item.setText(str(avg_price))
            avg_item.setTextAlignment(
                Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)

    def _recalc_parts_table(self, table):
        with blocked_signals(table):
            for r in range(table.rowCount()):
                self._recalc_parts_row(table, r)

    def _recalc_table(self, table):
        with blocked_signals(table):
            for r in range(table.rowCount()):
                self._recalc_row(table, r)

    def _recalc_all(self):
        with wait_cursor():
            self._recalc_table(self.works.findChild(QTableWidget))
            self._recalc_table(self.paint.findChild(QTableWidget))
            self._apply_wear_to_all_parts()
            for tab in [self.works, self.paint, self.mats, self.parts]:
                self._update_table_total(tab)
            self._update_total_summary()

    def _update_avg_total(self):
        if not hasattr(self, 'avg_prices_table') or not hasattr(self, 'avg_total_label'):
            return
        table = self.avg_prices_table
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        if "Средняя цена" not in headers:
            return
        avg_idx = headers.index("Средняя цена")
        total = count = 0
        for r in range(table.rowCount()):
            item = table.item(r, avg_idx)
            if item:
                try:
                    total += int(item.text().strip())
                    count += 1
                except (ValueError, TypeError):
                    pass
        total_str = f"{total:,}".replace(",", " ")
        self.avg_total_label.setText(
            f"💰 Средняя сумма: {total_str} ₽   ({count} поз.)")

    def _update_rate(self, val):
        try:
            v = float(val)
            if v > 0:
                self.rate = v
                self._recalc_all()
        except (ValueError, TypeError):
            pass

    def _update_rp_work(self, val):
        try:
            self.rp_work = int(val)
            self._recalc_all()
        except (ValueError, TypeError):
            pass

    def _update_rp_paint(self, val):
        try:
            self.rp_paint = int(val)
            self._recalc_all()
        except (ValueError, TypeError):
            pass

    def _update_wear(self, val):
        try:
            v = float(val)
            if 0 <= v <= 100:
                self.wear_percent = v
                self._apply_wear_to_all_parts()
                self._update_table_total(self.parts)
                self._update_total_summary()
        except (ValueError, TypeError):
            pass

    def _update_small_parts_percent(self, val):
        try:
            v = float(val)
            if 0 <= v <= 100:
                self.small_parts_percent = v
        except (ValueError, TypeError):
            pass

    def _apply_wear_to_all_parts(self):
        table = self.parts.findChild(QTableWidget)
        if not table:
            return
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        try:
            wear_idx = headers.index("Износ")
            cost_idx = headers.index("Стоимость")
            price_with_wear_idx = headers.index("Цена с износом")
        except ValueError:
            return
        with blocked_signals(table):
            for r in range(table.rowCount()):
                wear_item = table.item(r, wear_idx) or QTableWidgetItem()
                table.setItem(r, wear_idx, wear_item)
                wear_item.setText(str(self.wear_percent))
                cost_item = table.item(r, cost_idx)
                if cost_item:
                    try:
                        cost = float(cost_item.text().strip())
                        wear = max(0.0, min(100.0, self.wear_percent))
                        price_with_wear = round(cost * (1 - wear / 100))
                        pww_item = table.item(
                            r, price_with_wear_idx) or QTableWidgetItem()
                        table.setItem(r, price_with_wear_idx, pww_item)
                        pww_item.setText(str(price_with_wear))
                    except (ValueError, TypeError):
                        pass
            table.viewport().update()

    def _merge_continuations(self, table):
        with blocked_signals(table):
            merged = 0
            i = 1
            while i < table.rowCount():
                first = table.item(i, 0)
                if not first or not first.text().strip():
                    for c in range(1, table.columnCount()):
                        curr = table.item(i, c)
                        prev = table.item(i - 1, c)
                        if curr and curr.text().strip():
                            if prev:
                                prev.setText(prev.text() + " " +
                                             curr.text().strip())
                            else:
                                table.setItem(
                                    i - 1, c, QTableWidgetItem(curr.text().strip()))
                    table.removeRow(i)
                    merged += 1
                else:
                    i += 1
            table.viewport().update()
            if getattr(table, 'is_parts', False):
                self._apply_wear_to_all_parts()
            for tab in [self.works, self.paint, self.mats, self.parts]:
                if tab.findChild(QTableWidget) == table:
                    self._update_table_total(tab)
                    self._update_total_summary()
                    break
        QMessageBox.information(self, "Готово", f"Объединено: {merged}")

    def _apply_avg_prices(self):
        if not hasattr(self, 'avg_prices_table'):
            return
        avg_table = self.avg_prices_table
        parts_table = self.parts.findChild(QTableWidget)
        if avg_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Таблица средних цен пуста!")
            return
        if not parts_table:
            return
        avg_headers = [avg_table.horizontalHeaderItem(
            i).text() for i in range(avg_table.columnCount())]
        parts_headers = [parts_table.horizontalHeaderItem(
            i).text() for i in range(parts_table.columnCount())]
        if "Артикул" not in avg_headers or "Средняя цена" not in avg_headers:
            QMessageBox.critical(
                self, "Ошибка", "В таблице средних цен нет нужных колонок")
            return
        if "Артикул" not in parts_headers or "Стоимость" not in parts_headers:
            QMessageBox.critical(
                self, "Ошибка", "В таблице запчастей нет нужных колонок")
            return
        avg_art_idx = avg_headers.index("Артикул")
        avg_name_idx = avg_headers.index(
            "Наименование") if "Наименование" in avg_headers else -1
        avg_price_idx = avg_headers.index("Средняя цена")
        parts_art_idx = parts_headers.index("Артикул")
        parts_cost_idx = parts_headers.index("Стоимость")
        parts_name_idx = parts_headers.index(
            "Наименование") if "Наименование" in parts_headers else -1
        avg_data = []
        for r in range(avg_table.rowCount()):
            art_item = avg_table.item(r, avg_art_idx)
            price_item = avg_table.item(r, avg_price_idx)
            name_item = avg_table.item(
                r, avg_name_idx) if avg_name_idx >= 0 else None
            art = art_item.text().strip() if art_item else ""
            price = price_item.text().strip() if price_item else ""
            name = name_item.text().strip() if name_item else ""
            if not art or not price:
                continue
            try:
                price_val = int(price)
                if price_val > 0:
                    avg_data.append((art, name, price_val))
            except (ValueError, TypeError):
                continue
        if not avg_data:
            QMessageBox.warning(
                self, "Внимание", "Нет валидных данных для переноса!")
            return
        with wait_cursor(), blocked_signals(parts_table):
            updated = added = 0
            for art, name, price in avg_data:
                found_row = -1
                for r in range(parts_table.rowCount()):
                    item = parts_table.item(r, parts_art_idx)
                    if item and item.text().strip() == art:
                        found_row = r
                        break
                if found_row >= 0:
                    cost_item = parts_table.item(
                        found_row, parts_cost_idx) or QTableWidgetItem()
                    parts_table.setItem(found_row, parts_cost_idx, cost_item)
                    cost_item.setText(str(price))
                    updated += 1
                else:
                    new_row = parts_table.rowCount()
                    parts_table.insertRow(new_row)
                    parts_table.setItem(
                        new_row, parts_art_idx, QTableWidgetItem(art))
                    if parts_name_idx >= 0 and name:
                        parts_table.setItem(
                            new_row, parts_name_idx, QTableWidgetItem(name))
                    parts_table.setItem(
                        new_row, parts_cost_idx, QTableWidgetItem(str(price)))
                    added += 1
            self._apply_wear_to_all_parts()
            self._update_table_total(self.parts)
            self._update_total_summary()
            parts_table.viewport().update()
        QMessageBox.information(
            self, "✅ Готово",
            f"Средние цены перенесены в таблицу «Запчасти»:\n"
            f"• Обновлено существующих строк: {updated}\n"
            f"• Добавлено новых строк: {added}\n"
            f"• Всего обработано: {updated + added}\n"
            f"Сопоставление выполнено по артикулу."
        )

    def _transfer_parts_to_avg_prices(self):
        parts_table = self.parts.findChild(QTableWidget)
        if not parts_table or not hasattr(self, 'avg_prices_table'):
            return
        parts_headers = [parts_table.horizontalHeaderItem(
            i).text() for i in range(parts_table.columnCount())]
        avg_headers = [self.avg_prices_table.horizontalHeaderItem(
            i).text() for i in range(self.avg_prices_table.columnCount())]
        if "Артикул" not in parts_headers or "Наименование" not in parts_headers:
            QMessageBox.critical(
                self, "Ошибка", "В таблице запчастей нет нужных колонок")
            return
        if "Артикул" not in avg_headers or "Наименование" not in avg_headers:
            QMessageBox.critical(
                self, "Ошибка", "В таблице средних цен нет нужных колонок")
            return
        parts_art_idx = parts_headers.index("Артикул")
        parts_name_idx = parts_headers.index("Наименование")
        avg_art_idx = avg_headers.index("Артикул")
        avg_name_idx = avg_headers.index("Наименование")
        existing_articles = set()
        for r in range(self.avg_prices_table.rowCount()):
            item = self.avg_prices_table.item(r, avg_art_idx)
            if item and item.text().strip():
                existing_articles.add(item.text().strip())
        parts_data = []
        for r in range(parts_table.rowCount()):
            art_item = parts_table.item(r, parts_art_idx)
            name_item = parts_table.item(r, parts_name_idx)
            art = art_item.text().strip() if art_item else ""
            name = name_item.text().strip() if name_item else ""
            if art:
                parts_data.append((art, name))
        if not parts_data:
            QMessageBox.warning(
                self, "Внимание", "В таблице запчастей нет данных с артикулами!")
            return
        with wait_cursor(), blocked_signals(self.avg_prices_table):
            added = skipped = 0
            for art, name in parts_data:
                if art in existing_articles:
                    skipped += 1
                    continue
                new_row = self.avg_prices_table.rowCount()
                self.avg_prices_table.insertRow(new_row)
                self.avg_prices_table.setItem(
                    new_row, avg_art_idx, QTableWidgetItem(art))
                if name:
                    self.avg_prices_table.setItem(
                        new_row, avg_name_idx, QTableWidgetItem(name))
                existing_articles.add(art)
                added += 1
            self.avg_prices_table.viewport().update()
        QMessageBox.information(
            self, "✅ Готово",
            f"Перенос наименований и артикулов в таблицу средних цен:\n"
            f"• Добавлено новых строк: {added}\n"
            f"• Пропущено (уже есть): {skipped}\n"
            f"• Всего обработано: {added + skipped}\n"
            f"Теперь заполните цены в колонках «Источник 1-5», и нажмите «📥 Использовать средние цены на запчасти»."
        )

    def _read_file_with_encoding(self, path: str) -> str:
        with open(path, "rb") as f:
            raw_bytes = f.read()
        charset_match = re.search(
            rb'<meta[^>]+charset=["\']?([^"\'\s;>]+)', raw_bytes[:2000], re.IGNORECASE)
        if charset_match:
            declared_charset = charset_match.group(
                1).decode("ascii", errors="ignore").strip()
            try:
                return raw_bytes.decode(declared_charset)
            except (UnicodeDecodeError, LookupError):
                pass
        for encoding in ["utf-8", "cp1251", "windows-1251", "cp1252", "latin-1"]:
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("latin-1", errors="replace")

    def load_html(self):
        if not HAS_BS4:
            return QMessageBox.critical(self, "Ошибка", "Установите: pip install beautifulsoup4")
        path, _ = QFileDialog.getOpenFileName(
            self, "HTML", "", "HTML (*.html)")
        if not path:
            return
        with wait_cursor():
            try:
                html_content = self._read_file_with_encoding(path)
                # НОВОЕ: предобработка BRE Client HTML
                processed_text = AudatexParser.preprocess_html(html_content)
                self.editor.setPlainText(processed_text)
                self.status.setText(f"✅ Загружен: {os.path.basename(path)}")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", str(e))

    def _change_dict_path(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл словаря", "", "JSON файлы (*.json);;Все файлы (*.*)")
        if path:
            self.config_mgr.set_dict_path(path)
            self._manager = AbbrevManager(path)
            self._update_dict_path_label()
            self._refresh_all_tables()
            QMessageBox.information(
                self, "✅", f"Путь к базе данных изменён:\n{path}")

    def _update_dict_path_label(self):
        dict_path = self.config_mgr.get_dict_path()
        self.lbl_dict_path.setText(
            f"📁 База данных: {dict_path if dict_path else 'не выбрана'}")

    def open_dict_editor(self):
        if not self.manager or not self.manager.dict_path:
            self._change_dict_path()
            if not self.manager or not self.manager.dict_path:
                return
        dlg = QDialog(self)
        dlg.setWindowTitle("📖 Словарь")
        dlg.resize(600, 400)
        layout = QVBoxLayout(dlg)
        tree = QTreeWidget()
        tree.setHeaderLabels(["Сокращение", "Расшифровка"])
        tree.setColumnWidth(0, 200)
        for a, e in sorted(self.manager.get_all().items()):
            tree.addTopLevelItem(QTreeWidgetItem([a, e]))
        layout.addWidget(tree)

        def _add():
            a, _ = QInputDialog.getText(dlg, "Сокращение", "Текст:")
            e, _ = QInputDialog.getText(dlg, "Расшифровка", "Текст:")
            if a and e:
                self.manager.add(a, e)
                tree.addTopLevelItem(QTreeWidgetItem([a.upper(), e]))
                self._refresh_all_tables()

        def _rem():
            sel = tree.currentItem()
            if sel:
                self.manager.remove(sel.text(0))
                tree.takeTopLevelItem(tree.indexOfTopLevelItem(sel))
                self._refresh_all_tables()

        def _clear():
            if QMessageBox.question(self, "Очистить?", "Удалить все сокращения?") == QMessageBox.StandardButton.Yes:
                self.manager.abbreviations.clear()
                self.manager.save()
                tree.clear()
                self._refresh_all_tables()
        btns = QHBoxLayout()
        for txt, cmd in [("Добавить", _add), ("Удалить", _rem), ("Очистить всё", _clear), ("Закрыть", dlg.accept)]:
            btn = create_styled_button(txt, *BTN_LIGHT, padding=4)
            btn.clicked.connect(cmd)
            btns.addWidget(btn)
        btns.addStretch()
        layout.addLayout(btns)
        dlg.exec()

    def _refresh_all_tables(self):
        for tab in [self.works, self.paint, self.mats, self.parts]:
            t = tab.findChild(QTableWidget)
            if t:
                t.setItemDelegate(AbbrevDelegate(t, self._manager, getattr(
                    t, 'highlight_column', "Описание"), self.is_dark_mode))
                t.viewport().update()

    def _save_calculation(self):
        path = self.last_saved_path
        if not path:
            path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить расчёт", "calculation.json", "JSON файлы (*.json);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                data = {
                    "version": CALC_VERSION,
                    "params": {"rate": self.rate, "rp_work": self.rp_work, "rp_paint": self.rp_paint, "wear_percent": self.wear_percent, "small_parts_percent": self.small_parts_percent},
                    "tables": {
                        "works": self._get_table_data(self.works.findChild(QTableWidget)),
                        "paint": self._get_table_data(self.paint.findChild(QTableWidget)),
                        "mats": self._get_table_data(self.mats.findChild(QTableWidget)),
                        "parts": self._get_table_data(self.parts.findChild(QTableWidget)),
                        "avg_prices": self._get_table_data(self.avg_prices_table) if hasattr(self, 'avg_prices_table') else [],
                    },
                    "avg_sources": [edit.text() for edit in self.source_edits] if hasattr(self, 'source_edits') else [],
                    "dict_path": self.manager.dict_path if self.manager else "",
                }
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
                self.last_saved_path = path
                self.setWindowTitle(
                    f"AUDATEX Converter v{CALC_VERSION} — {os.path.basename(path)}")
                self.status.setText(
                    f"💾 Расчёт сохранён: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "✅ Сохранено", f"Расчёт сохранён:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось сохранить расчёт:\n{e}")

    def _load_calculation(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Загрузить расчёт", "", "JSON файлы (*.json);;Все файлы (*.*)")
        if not path:
            return
        with wait_cursor():
            try:
                with open(path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                version = data.get("version", "unknown")
                if version != CALC_VERSION:
                    reply = QMessageBox.question(self, "Несовместимая версия", f"Версия файла: {version}\nВерсия программы: {CALC_VERSION}\nПродолжить загрузку?",
                                                 QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
                    if reply == QMessageBox.StandardButton.No:
                        return
                params = data.get("params", {})
                self.rate = params.get("rate", 285.0)
                self.rp_work = params.get("rp_work", 10)
                self.rp_paint = params.get("rp_paint", 10)
                self.wear_percent = params.get("wear_percent", 0.0)
                self.small_parts_percent = params.get(
                    "small_parts_percent", 0.0)
                with blocked_signals(self.in_rate, self.in_wear, self.in_small_parts, self.in_rp_w, self.in_rp_p):
                    for widget, value in [(self.in_rate, str(int(self.rate))), (self.in_wear, str(int(self.wear_percent))), (self.in_small_parts, str(int(self.small_parts_percent)))]:
                        widget.setText(value)
                    for combo, value in [(self.in_rp_w, str(self.rp_work)), (self.in_rp_p, str(self.rp_paint))]:
                        combo.setCurrentText(value)
                dict_path = data.get("dict_path", "")
                if dict_path and os.path.exists(dict_path):
                    self.config_mgr.set_dict_path(dict_path)
                    self._manager = AbbrevManager(dict_path)
                    self._update_dict_path_label()
                tables = data.get("tables", {})
                self._set_table_data(self.works.findChild(
                    QTableWidget), tables.get("works", []))
                self._set_table_data(self.paint.findChild(
                    QTableWidget), tables.get("paint", []))
                self._set_table_data(self.mats.findChild(
                    QTableWidget), tables.get("mats", []))
                self._set_table_data(self.parts.findChild(
                    QTableWidget), tables.get("parts", []))
                if hasattr(self, 'avg_prices_table'):
                    self._set_table_data(
                        self.avg_prices_table, tables.get("avg_prices", []))
                sources = data.get("avg_sources", [])
                if hasattr(self, 'source_edits'):
                    for i, edit in enumerate(self.source_edits):
                        edit.setText(sources[i] if i < len(sources) else "")
                self._apply_wear_to_all_parts()
                self._recalc_all()
                self.last_saved_path = path
                self.setWindowTitle(
                    f"AUDATEX Converter v{CALC_VERSION} — {os.path.basename(path)}")
                self.status.setText(
                    f"📂 Расчёт загружен: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "✅ Загружено", f"Расчёт загружен:\n{path}")
            except Exception as e:
                QMessageBox.critical(
                    self, "Ошибка", f"Не удалось загрузить расчёт:\n{e}")

    def _get_table_data(self, table: QTableWidget) -> List[List[str]]:
        if not table:
            return []
        data = []
        for r in range(table.rowCount()):
            row = []
            for c in range(table.columnCount()):
                item = table.item(r, c)
                row.append(item.text() if item else "")
            data.append(row)
        return data

    def _set_table_data(self, table: QTableWidget, data: List[List[str]]):
        if not table:
            return
        with blocked_signals(table):
            table.setRowCount(0)
            for row in data:
                r = table.rowCount()
                table.insertRow(r)
                for c, val in enumerate(row):
                    if c < table.columnCount():
                        table.setItem(r, c, QTableWidgetItem(val))
            table.viewport().update()

    def _get_parts_total_without_wear(self) -> int:
        table = self.parts.findChild(QTableWidget)
        if not table:
            return 0
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        if "Стоимость" not in headers:
            return 0
        cost_idx = headers.index("Стоимость")
        total = 0
        for r in range(table.rowCount()):
            item = table.item(r, cost_idx)
            if item:
                try:
                    total += int(item.text().strip())
                except (ValueError, TypeError):
                    pass
        return total

    def _get_parts_total_with_wear(self) -> int:
        table = self.parts.findChild(QTableWidget)
        if not table:
            return 0
        headers = [table.horizontalHeaderItem(
            i).text() for i in range(table.columnCount())]
        if "Цена с износом" not in headers:
            return 0
        cost_idx = headers.index("Цена с износом")
        total = 0
        for r in range(table.rowCount()):
            item = table.item(r, cost_idx)
            if item:
                try:
                    total += int(item.text().strip())
                except (ValueError, TypeError):
                    pass
        return total

    def _add_small_parts_row(self, headers: List[str], data: List[List[str]]) -> List[List[str]]:
        if self.small_parts_percent <= 0:
            return data
        parts_total = self._get_parts_total_without_wear()
        small_parts_cost = round(parts_total * self.small_parts_percent / 100)
        if small_parts_cost <= 0:
            return data
        description = f"Мелкие запчасти {int(self.small_parts_percent)}% от стоимости запасных частей без износа"
        new_row = ["" for _ in headers]
        if "Описание" in headers:
            new_row[headers.index("Описание")] = description
        if "Стоимость" in headers:
            new_row[headers.index("Стоимость")] = str(small_parts_cost)
        result = list(data)
        result.append(new_row)
        return result

    def _collect_tables_data(self) -> List[tuple]:
        tables_data = []
        for name, tab in [("Работы", self.works), ("Окраска", self.paint), ("Материалы", self.mats), ("Запчасти", self.parts)]:
            t = tab.findChild(QTableWidget)
            if t.rowCount() == 0 and name != "Материалы":
                if self.small_parts_percent <= 0:
                    continue
            headers = [t.horizontalHeaderItem(
                c).text() for c in range(t.columnCount())]
            data = []
            for r in range(t.rowCount()):
                row = []
                for c in range(t.columnCount()):
                    val = t.item(r, c).text() if t.item(r, c) else ""
                    if self.chk_expand.isChecked() and headers[c] in ["Описание", "Наименование"]:
                        val = self.manager.expand(val) if self.manager else val
                    row.append(val)
                data.append(row)
            if name == "Материалы":
                data = self._add_small_parts_row(headers, data)
            if name == "Запчасти" and data:
                total_without_wear = self._get_parts_total_without_wear()
                total_with_wear = self._get_parts_total_with_wear()
                data.append(["" for _ in headers])
                total_row = ["" for _ in headers]
                if "Наименование" in headers:
                    total_row[headers.index("Наименование")] = "ИТОГО:"
                if "Стоимость" in headers:
                    total_row[headers.index("Стоимость")] = str(
                        total_without_wear)
                if "Цена с износом" in headers:
                    total_row[headers.index("Цена с износом")] = str(
                        total_with_wear)
                data.append(total_row)
            if not data:
                continue
            tables_data.append(
                (name, headers, data, getattr(t, 'is_parts', False)))
        return tables_data

    def _is_total_row(self, name: str, row_idx: int, row: List[str], headers: List[str]) -> bool:
        for col_name in ["Наименование", "Описание"]:
            if col_name in headers:
                col_idx = headers.index(col_name)
                if col_idx < len(row) and row[col_idx] and "ИТОГО" in row[col_idx].upper():
                    return True
        return False

    def _get_cost_column_info(self, headers: List[str], is_parts: bool) -> Tuple[Optional[str], int]:
        if is_parts and "Цена с износом" in headers:
            cost_col_name = "Цена с износом"
        elif "Стоимость" in headers:
            cost_col_name = "Стоимость"
        else:
            return None, -1
        return cost_col_name, headers.index(cost_col_name)

    def export_ods(self):
        if not HAS_ODF:
            return QMessageBox.critical(self, "Ошибка", "Установите: pip install odfpy")
        tables_data = self._collect_tables_data()
        if not tables_data:
            return QMessageBox.warning(self, "Внимание", "Нет данных для экспорта!")
        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить", "AUDATEX_export.ods", "ODS (*.ods)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = OpenDocumentSpreadsheet()
                styles = {"TS_TNR12": self._make_ods_style("TS_TNR12", bold=False, align_right=False), "TS_TNR12_Right": self._make_ods_style("TS_TNR12_Right", bold=False, align_right=True), "TS_TNR12_Bold": self._make_ods_style(
                    "TS_TNR12_Bold", bold=True, align_right=False), "TS_TNR12_Bold_Right": self._make_ods_style("TS_TNR12_Bold_Right", bold=True, align_right=True)}
                for style in styles.values():
                    doc.automaticstyles.addElement(style)
                for name, headers, data, is_parts in tables_data:
                    tbl = Table(name=name)
                    for _ in headers:
                        tbl.addElement(TableColumn())
                    r = TableRow()
                    for h in headers:
                        c = TableCell(valuetype="string",
                                      stylename="TS_TNR12_Bold")
                        c.addElement(P(text=h))
                        r.addElement(c)
                    tbl.addElement(r)
                    _, cost_idx = self._get_cost_column_info(headers, is_parts)
                    for row_idx, row in enumerate(data):
                        is_total = self._is_total_row(
                            name, row_idx, row, headers)
                        r = TableRow()
                        for i, val in enumerate(row):
                            header = headers[i]
                            if is_total:
                                style = "TS_TNR12_Bold_Right" if is_numeric_column(
                                    header) else "TS_TNR12_Bold"
                            elif is_numeric_column(header):
                                style = "TS_TNR12_Right"
                            else:
                                style = "TS_TNR12"
                            c = TableCell(valuetype="string", stylename=style)
                            c.addElement(P(text=str(val)))
                            r.addElement(c)
                        tbl.addElement(r)
                    if cost_idx >= 0 and name != "Запчасти":
                        total = sum(int(row[cost_idx]) for row in data if not self._is_total_row(name, data.index(
                            row), row, headers) and cost_idx < len(row) and row[cost_idx].strip().isdigit())
                        tbl.addElement(TableRow())
                        r = TableRow()
                        for i, h in enumerate(headers):
                            style = "TS_TNR12_Bold_Right" if i == cost_idx else "TS_TNR12_Bold"
                            c = TableCell(valuetype="string", stylename=style)
                            txt = "ИТОГО:" if i == 0 else (
                                str(total) if i == cost_idx else "")
                            c.addElement(P(text=txt))
                            r.addElement(c)
                        tbl.addElement(r)
                    doc.spreadsheet.addElement(tbl)
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортировано: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Файл сохранён:\n{path}")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", str(e))

    def _make_ods_style(self, name: str, bold: bool, align_right: bool):
        style = Style(name=name, family="table-cell")
        style.addElement(TextProperties(
            fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt", fontweight="bold" if bold else "normal"))
        style.addElement(ParagraphProperties(
            lineheight="100%", textalign="end" if align_right else "start"))
        return style

    def export_odt(self):
        if not HAS_ODF:
            return QMessageBox.critical(self, "Ошибка", "Установите: pip install odfpy")
        tables_data = self._collect_tables_data()
        if not tables_data:
            return QMessageBox.warning(self, "Внимание", "Нет данных для экспорта!")
        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить", "AUDATEX_export.odt", "ODT (*.odt)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = OpenDocumentText()
                title_style = Style(name="Title", family="paragraph")
                title_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize="18pt", fontweight="bold"))
                title_style.addElement(ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(title_style)
                heading_style = Style(
                    name="Heading_Custom", family="paragraph")
                heading_style.addElement(TextProperties(
                    fontname=FONT_EXPORT, fontsize="14pt", fontweight="bold"))
                heading_style.addElement(
                    ParagraphProperties(lineheight="100%"))
                doc.automaticstyles.addElement(heading_style)
                for style_name, bold, align_right in [("Cell_TNR12", False, False), ("Cell_TNR12_Right", False, True), ("Cell_TNR12_Bold", True, False), ("Cell_TNR12_Bold_Right", True, True)]:
                    style = Style(name=style_name, family="table-cell")
                    style.addElement(TextProperties(
                        fontname=FONT_EXPORT, fontsize=f"{FONT_EXPORT_SIZE}pt", fontweight="bold" if bold else "normal"))
                    style.addElement(ParagraphProperties(
                        lineheight="100%", textalign="end" if align_right else "start"))
                    doc.automaticstyles.addElement(style)
                doc.text.addElement(P(text="Смета AUDATEX", stylename="Title"))
                for name, headers, data, is_parts in tables_data:
                    doc.text.addElement(
                        P(text=f"\n{name}", stylename="Heading_Custom"))
                    tbl = Table(name=name)
                    for _ in headers:
                        tbl.addElement(TableColumn())
                    r = TableRow()
                    for h in headers:
                        c = TableCell(valuetype="string",
                                      stylename="Cell_TNR12_Bold")
                        c.addElement(P(text=h))
                        r.addElement(c)
                    tbl.addElement(r)
                    _, cost_idx = self._get_cost_column_info(headers, is_parts)
                    for row_idx, row in enumerate(data):
                        is_total = self._is_total_row(
                            name, row_idx, row, headers)
                        r = TableRow()
                        for i, val in enumerate(row):
                            header = headers[i]
                            if is_total:
                                style = "Cell_TNR12_Bold_Right" if is_numeric_column(
                                    header) else "Cell_TNR12_Bold"
                            elif is_numeric_column(header):
                                style = "Cell_TNR12_Right"
                            else:
                                style = "Cell_TNR12"
                            c = TableCell(valuetype="string", stylename=style)
                            c.addElement(P(text=str(val)))
                            r.addElement(c)
                        tbl.addElement(r)
                    if cost_idx >= 0 and name != "Запчасти":
                        total = sum(int(row[cost_idx]) for row in data if not self._is_total_row(name, data.index(
                            row), row, headers) and cost_idx < len(row) and row[cost_idx].strip().isdigit())
                        tbl.addElement(TableRow())
                        r = TableRow()
                        for i, h in enumerate(headers):
                            style = "Cell_TNR12_Bold_Right" if i == cost_idx else "Cell_TNR12_Bold"
                            c = TableCell(valuetype="string", stylename=style)
                            txt = "ИТОГО:" if i == 0 else (
                                str(total) if i == cost_idx else "")
                            c.addElement(P(text=txt))
                            r.addElement(c)
                        tbl.addElement(r)
                    doc.text.addElement(tbl)
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортировано: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Файл сохранён:\n{path}")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", str(e))

    def export_excel(self):
        if not HAS_XLSX:
            return QMessageBox.critical(self, "Ошибка", "Установите: pip install openpyxl")
        tables_data = self._collect_tables_data()
        if not tables_data:
            return QMessageBox.warning(self, "Внимание", "Нет данных для экспорта!")
        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить", "AUDATEX_export.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        with wait_cursor():
            try:
                wb = Workbook()
                wb.remove(wb.active)
                header_font = XlFont(
                    name=FONT_EXPORT, size=FONT_EXPORT_SIZE, bold=True, color="FFFFFF")
                header_fill = PatternFill(
                    start_color="2196F3", end_color="2196F3", fill_type="solid")
                total_font = XlFont(
                    name=FONT_EXPORT, size=FONT_EXPORT_SIZE, bold=True, color="FFFFFF")
                total_fill = PatternFill(
                    start_color="43A047", end_color="43A047", fill_type="solid")
                border = Border(left=Side(style='thin'), right=Side(
                    style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                align_left = Alignment(
                    horizontal='left', vertical='center', wrap_text=True)
                align_right = Alignment(horizontal='right', vertical='center')
                for name, headers, data, is_parts in tables_data:
                    ws = wb.create_sheet(title=name[:31])
                    for col_idx, h in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=col_idx, value=h)
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.border = border
                        cell.alignment = align_left
                    _, cost_idx = self._get_cost_column_info(headers, is_parts)
                    for row_idx, row in enumerate(data, 2):
                        is_total = self._is_total_row(
                            name, row_idx - 2, row, headers)
                        for col_idx, val in enumerate(row, 1):
                            header = headers[col_idx - 1]
                            cell = ws.cell(row=row_idx, column=col_idx)
                            if is_numeric_column(header):
                                try:
                                    val = int(val)
                                except (ValueError, TypeError):
                                    try:
                                        val = float(val)
                                    except (ValueError, TypeError):
                                        pass
                            cell.value = val
                            cell.font = XlFont(
                                name=FONT_EXPORT, size=FONT_EXPORT_SIZE, bold=is_total)
                            cell.border = border
                            cell.alignment = align_right if is_numeric_column(
                                header) else align_left
                    if cost_idx >= 0 and name != "Запчасти":
                        total = sum(int(row[cost_idx]) for row in data if cost_idx < len(
                            row) and row[cost_idx].strip().isdigit())
                        total_row = ws.max_row + 2
                        ws.cell(row=total_row, column=1,
                                value="ИТОГО:").font = total_font
                        ws.cell(row=total_row, column=1).fill = total_fill
                        ws.cell(row=total_row, column=1).border = border
                        ws.cell(row=total_row, column=1).alignment = align_left
                        total_cell = ws.cell(
                            row=total_row, column=cost_idx + 1, value=total)
                        total_cell.font = total_font
                        total_cell.fill = total_fill
                        total_cell.border = border
                        total_cell.alignment = align_right
                        total_cell.number_format = '#,##0 ₽'
                    for col in ws.columns:
                        max_length = max((len(str(cell.value))
                                         for cell in col if cell.value), default=0)
                        ws.column_dimensions[col[0].column_letter].width = min(
                            max_length + 2, 50)
                wb.save(path)
                self.status.setText(
                    f"✅ Экспортировано: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Файл сохранён:\n{path}")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", str(e))

    def export_word(self):
        if not HAS_DOCX:
            return QMessageBox.critical(self, "Ошибка", "Установите: pip install python-docx")
        tables_data = self._collect_tables_data()
        if not tables_data:
            return QMessageBox.warning(self, "Внимание", "Нет данных для экспорта!")
        path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить", "AUDATEX_export.docx", "Word (*.docx)")
        if not path:
            return
        with wait_cursor():
            try:
                doc = DocxDocument()
                title = doc.add_heading('Смета AUDATEX', level=0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._apply_docx_font(title, size=18, bold=True)
                for name, headers, data, is_parts in tables_data:
                    heading = doc.add_heading(name, level=1)
                    self._apply_docx_font(heading, size=14, bold=True)
                    table = doc.add_table(rows=1, cols=len(headers))
                    table.style = 'Table Grid'
                    hdr_cells = table.rows[0].cells
                    for i, h in enumerate(headers):
                        hdr_cells[i].text = h
                        for paragraph in hdr_cells[i].paragraphs:
                            self._apply_docx_paragraph_format(paragraph)
                            for run in paragraph.runs:
                                self._apply_docx_run_font(run, bold=True)
                    _, cost_idx = self._get_cost_column_info(headers, is_parts)
                    for row in data:
                        row_cells = table.add_row().cells
                        is_total = self._is_total_row(
                            name, data.index(row), row, headers)
                        for i, val in enumerate(row):
                            row_cells[i].text = str(val)
                            header = headers[i]
                            for paragraph in row_cells[i].paragraphs:
                                self._apply_docx_paragraph_format(paragraph)
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if is_numeric_column(
                                    header) else WD_ALIGN_PARAGRAPH.LEFT
                                for run in paragraph.runs:
                                    self._apply_docx_run_font(
                                        run, bold=is_total)
                    if cost_idx >= 0 and name != "Запчасти":
                        total = sum(int(row[cost_idx]) for row in data if cost_idx < len(
                            row) and row[cost_idx].strip().isdigit())
                        total_row = table.add_row().cells
                        total_row[0].text = "ИТОГО:"
                        total_row[cost_idx].text = f"{total:,}".replace(
                            ",", " ") + " ₽"
                        for cell in total_row:
                            for paragraph in cell.paragraphs:
                                self._apply_docx_paragraph_format(paragraph)
                                for run in paragraph.runs:
                                    self._apply_docx_run_font(run, bold=True)
                        for paragraph in total_row[cost_idx].paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    doc.add_paragraph()
                doc.save(path)
                self.status.setText(
                    f"✅ Экспортировано: {os.path.basename(path)}")
                QMessageBox.information(
                    self, "Успех", f"Файл сохранён:\n{path}")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", str(e))

    def _apply_docx_font(self, heading, size: int, bold: bool):
        for run in heading.runs:
            run.font.name = FONT_EXPORT
            run.font.size = Pt(size)
            run.font.bold = bold
            r = run._element
            rPr = r.get_or_add_rPr()
            from lxml import etree
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = etree.SubElement(rPr, qn('w:rFonts'))
            rFonts.set(qn('w:eastAsia'), FONT_EXPORT)

    def _apply_docx_paragraph_format(self, paragraph):
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

    def _apply_docx_run_font(self, run, bold: bool = False):
        run.font.name = FONT_EXPORT
        run.font.size = Pt(FONT_EXPORT_SIZE)
        run.font.bold = bold
        r = run._element
        rPr = r.get_or_add_rPr()
        from lxml import etree
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = etree.SubElement(rPr, qn('w:rFonts'))
        rFonts.set(qn('w:eastAsia'), FONT_EXPORT)


def main():
    app = QApplication(sys.argv)
    if getattr(sys, 'frozen', False):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    splash_image_path = os.path.join(base_path, "splash.png")
    splash = SplashScreen(splash_image_path)
    splash.show()
    app.processEvents()
    messages = ["Инициализация модулей...", "Загрузка словарей...",
                "Подготовка интерфейса...", "Проверка обновлений...", "Готово к работе!"]
    for i in range(101):
        time.sleep(0.015)
        if i % 20 == 0:
            msg_idx = min(i // 20, len(messages) - 1)
            splash.update_progress(i, messages[msg_idx])
        else:
            splash.update_progress(i)
        app.processEvents()
    splash.close()
    app.processEvents()
    w = MainWindow()
    w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
